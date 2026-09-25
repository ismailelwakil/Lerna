"""Document pipeline (spec #10, #11): extraction → normalization → chunking →
metadata (page/slide/section) → malware-scan hook → storage.

Local libraries only (pypdf / python-docx / python-pptx); images go to the
Gemini vision model for OCR when configured."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Optional, Protocol

from ...core.logging import get_logger
from ...core.security import new_id

log = get_logger("docs")

CHUNK_DEFAULT = 1200
OVERLAP_DEFAULT = 150


@dataclass
class ExtractedChunk:
    text: str
    page: Optional[int] = None
    slide: Optional[int] = None
    section: Optional[str] = None


@dataclass
class ExtractionResult:
    pages: int = 0
    slides: int = 0
    chunks: list[ExtractedChunk] = field(default_factory=list)
    language_hint: str = "en"
    meta: dict = field(default_factory=dict)


class MalwareScanner(Protocol):
    def scan(self, content: bytes, filename: str) -> tuple[bool, str]:
        """Returns (clean, note). Hook point for ClamAV etc."""


class NullScanner:
    """Default dev hook — NO real malware protection (documented honestly).

    Production should set MALWARE_SCANNER=clamav (and typically
    REQUIRE_MALWARE_SCANNING=true) — see SECURITY.md."""
    name = "null"

    def scan(self, content: bytes, filename: str) -> tuple[bool, str]:
        return True, "no-scanner-configured"


class ClamAVScanner:
    """Real scanning via ClamAV's documented INSTREAM protocol (clamd TCP).

    No third-party dependency: raw socket, chunked INSTREAM, terminate with
    a zero-length chunk, read verdict ("stream: OK" / "stream: EICAR-... FOUND").
    """
    name = "clamav"

    def __init__(self, host: str, port: int, timeout: float = 30.0) -> None:
        self._host, self._port, self._timeout = host, port, timeout

    def scan(self, content: bytes, filename: str) -> tuple[bool, str]:
        import socket
        try:
            with socket.create_connection((self._host, self._port), timeout=self._timeout) as sock:
                sock.sendall(b"zINSTREAM\0")
                for offset in range(0, len(content), 32768):
                    chunk = content[offset:offset + 32768]
                    sock.sendall(len(chunk).to_bytes(4, "big") + chunk)
                sock.sendall((0).to_bytes(4, "big"))
                verdict = b""
                while True:
                    part = sock.recv(4096)
                    if not part:
                        break
                    verdict += part
                    if b"\0" in part or b"OK" in part or b"FOUND" in part:
                        break
        except OSError as exc:
            return False, f"scanner-unreachable: {type(exc).__name__}"
        text = verdict.decode("utf-8", "replace").strip("\0 \r\n")
        if text.endswith("OK"):
            return True, "clean"
        return False, f"infected: {text[:120]}"


def build_scanner():
    """Select the scanner from configuration (fail-safe in production)."""
    from ...core.config import get_settings
    from ...core.logging import get_logger
    settings = get_settings()
    log = get_logger("docs")
    if settings.malware_scanner == "clamav":
        log.info("malware_scanner_selected", extra={"ctx": {"scanner": "clamav",
                                                            "required": settings.require_malware_scanning}})
        return ClamAVScanner(settings.clamav_host, settings.clamav_port)
    if settings.require_malware_scanning:
        # production demands scanning but no scanner is configured → refuse
        # to run as if files were clean (fail-safe).
        log.error("malware_scanner_required_but_missing")
        return None
    return NullScanner()


scanner: MalwareScanner = NullScanner()  # replaced at pipeline runtime


# ------------------------------------------------------------------ extract
def extract_pdf(content: bytes) -> ExtractionResult:
    import io
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    result = ExtractionResult(pages=len(reader.pages))
    info = (reader.metadata or {})
    result.meta["title"] = str(info.title) if getattr(info, "title", None) else ""
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 — skip unreadable pages
            continue
        result.chunks.extend(_chunk(text, page=number))
    return result


def extract_docx(content: bytes) -> ExtractionResult:
    import io
    from docx import Document
    document = Document(io.BytesIO(content))
    result = ExtractionResult()
    section = None
    buffer: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            if buffer:
                result.chunks.extend(_chunk(" ".join(buffer), section=section))
                buffer = []
            section = text[:200]
        buffer.append(text)
    if buffer:
        result.chunks.extend(_chunk(" ".join(buffer), section=section))
    result.meta["headings"] = True
    return result


def extract_pptx(content: bytes) -> ExtractionResult:
    import io
    from pptx import Presentation
    presentation = Presentation(io.BytesIO(content))
    result = ExtractionResult(slides=len(presentation.slides))
    for index, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in paragraph.runs).strip()
                    if line:
                        parts.append(line)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    parts.append(" | ".join(cells))
        if parts:
            result.chunks.extend(_chunk("\n".join(parts), slide=index))
    return result


def extract_txt(content: bytes) -> ExtractionResult:
    text = content.decode("utf-8", "replace")
    result = ExtractionResult(pages=max(1, text.count("\f") + 1))
    result.chunks.extend(_chunk(text))
    return result


async def extract_image(content: bytes, mime: str) -> ExtractionResult:
    """OCR via Gemini vision when configured; honest failure otherwise."""
    from ...core.config import get_settings
    s = get_settings()
    result = ExtractionResult(pages=1)
    if not s.gemini_api_key:
        result.meta["ocr"] = "unavailable (no vision provider configured)"
        return result
    import base64
    import httpx
    try:
        async with httpx.AsyncClient(timeout=90.0) as client:
            resp = await client.post(
                f"{s.gemini_base}/models/{s.gemini_image_model}:generateContent",
                headers={"x-goog-api-key": s.gemini_api_key},
                json={"contents": [{"parts": [
                    {"text": "Extract ALL text from this image exactly. Output only the text."},
                    {"inlineData": {"mimeType": mime, "data": base64.b64encode(content).decode()}}]}]})
        if resp.status_code == 200:
            parts = resp.json()["candidates"][0]["content"]["parts"]
            text = "\n".join(p.get("text", "") for p in parts).strip()
            result.chunks.extend(_chunk(text))
            result.meta["ocr"] = "gemini"
            return result
        result.meta["ocr"] = f"failed ({resp.status_code})"
    except httpx.HTTPError as exc:
        result.meta["ocr"] = f"failed ({str(exc)[:80]})"
    return result


# ------------------------------------------------------------------ chunk
def _chunk(text: str, *, page: Optional[int] = None, slide: Optional[int] = None,
           section: Optional[str] = None, size: int = CHUNK_DEFAULT,
           overlap: int = OVERLAP_DEFAULT) -> list[ExtractedChunk]:
    text = re.sub(r"[ \t]+", " ", text or "").strip()
    if not text:
        return []
    out: list[ExtractedChunk] = []
    start = 0
    while start < len(text):
        piece = text[start:start + size]
        if len(piece) == size:  # break on word boundary
            cut = piece.rfind(" ")
            if cut > size * 0.6:
                piece = piece[:cut]
        out.append(ExtractedChunk(text=piece.strip(), page=page, slide=slide,
                                  section=section))
        if start + size >= len(text):
            break
        start += len(piece) - overlap
    return out


ARABIC = re.compile(r"[\u0600-\u06FF]")


def process(content: bytes, ext: str, mime: str = "") -> ExtractionResult:
    """Synchronous extraction for pdf/docx/pptx/txt; images handled async."""
    if ext == ".pdf":
        return extract_pdf(content)
    if ext == ".docx":
        return extract_docx(content)
    if ext == ".pptx":
        return extract_pptx(content)
    if ext in {".txt", ".md", ".csv"}:
        return extract_txt(content)
    raise ValueError(f"async extraction required for {ext}")


def detect_language(chunks: list[ExtractedChunk]) -> str:
    blob = " ".join(c.text for c in chunks[:20])
    return "ar" if ARABIC.search(blob) else "en"


def new_document_id() -> str:
    return new_id()