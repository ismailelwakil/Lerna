from __future__ import annotations

import json
import re
from pathlib import Path
import pytest

from src.contracts.models import Evidence
from src.orchestration.factory import build_system
from src.integration.service import build_module
from src.content_generation.service import ContentGenerator, ResourceContractValidator, ExamQuestion


def test_generator_consumes_learner_state_context(tmp_path):
    """1. Generator consumes learner-state context from the student profile."""
    ai = build_system(str(tmp_path), test_mode=True)
    p = ai.create_or_load_student("s-101", "Learner", "Computer Science", "en", "step-by-step")
    p.recent_topics = ["Binary search"]
    p.concept_mastery = {"Binary search": 1.0, "worst-case complexity": 0.0, "space complexity": 0.4}
    p.strengths = ["Binary search"]
    p.unknown_concepts = ["worst-case complexity"]
    p.weak_concepts = ["space complexity"]
    ai.store.save(p)

    out, _ = ai.generate_content("s-101", "worst-case complexity", ["explanation"], "en")
    assert "explanation" in out
    res = out["explanation"]
    assert res.get("learner_state") == "missing_knowledge"
    assert res.get("strategy") == "foundational"

    saved = ai.get_student_profile("s-101")
    last_gen = saved.generated_resources[-1]
    assert last_gen["learner_snapshot"]["unknown_concepts"] == ["worst-case complexity"]
    assert last_gen["learner_snapshot"]["weak_areas"] == ["space complexity"]


def test_missing_knowledge_gets_foundational_content(tmp_path):
    """2. Missing knowledge gets foundational explanation, basic terms, and no misconception blame."""
    ai = build_system(str(tmp_path), test_mode=True)
    p = ai.create_or_load_student("s-102", "Learner", "Computer Science", "en")
    p.recent_topics = ["Binary search"]
    p.unknown_concepts = ["worst-case complexity"]
    p.concept_mastery = {"worst-case complexity": 0.0}
    ai.store.save(p)

    out, _ = ai.generate_content("s-102", "worst-case complexity", ["explanation", "study_guide"], "en")
    exp = out["explanation"]["content"]
    assert "Foundational" in exp or "Basics" in exp
    assert "misconception" not in exp.lower() or "active — no prior mastery" in exp.lower()
    assert out["explanation"].get("strategy") == "foundational"


def test_weak_attempted_concept_gets_reinforcement_content(tmp_path):
    """3. Weak attempted concept gets reinforcement content, error contrast, and no prerequisite blame."""
    ai = build_system(str(tmp_path), test_mode=True)
    p = ai.create_or_load_student("s-103", "Learner", "Computer Science", "en")
    p.recent_topics = ["Binary search"]
    p.weak_concepts = ["space complexity"]
    p.unknown_concepts = []  # Not missing knowledge
    p.concept_mastery = {"space complexity": 0.35}
    ai.store.save(p)

    out, _ = ai.generate_content("s-103", "space complexity", ["explanation"], "en")
    exp = out["explanation"]["content"]
    assert "Reinforcement" in exp or "Reasoning" in exp
    assert "prerequisite gap" not in exp.lower() or "without prerequisite gap" in exp.lower()
    assert out["explanation"].get("strategy") == "reinforcement"


def test_strength_is_not_targeted_for_basic_remediation(tmp_path):
    """4. Strength is not targeted for basic remediation; advances difficulty using scaffolding."""
    ai = build_system(str(tmp_path), test_mode=True)
    p = ai.create_or_load_student("s-104", "Learner", "Computer Science", "en")
    p.strengths = ["Binary search"]
    p.concept_mastery = {"Binary search": 1.0}
    ai.store.save(p)

    out, _ = ai.generate_content("s-104", "Binary search", ["explanation"], "en")
    exp = out["explanation"]["content"]
    assert "Advanced" in exp or "Scaffolding" in exp or "In-Depth" in exp
    assert out["explanation"].get("strategy") == "higher-difficulty"


def test_selected_material_uses_uploaded_evidence_only(tmp_path):
    """5. When material is explicitly selected, generated content uses uploaded evidence only."""
    ai = build_system(str(tmp_path), test_mode=True)
    ai.create_or_load_student("s-105", "Learner", "Computer Science", "en")

    lecture = tmp_path / "sample_lecture.txt"
    lecture.write_text(
        "Binary search requires a sorted indexed array. It repeatedly halves the search space. "
        "Worst-case runtime is O(log n). Iterative space complexity is O(1).",
        encoding="utf-8",
    )
    doc_id = ai.upload_document("s-105", lecture, "sample_lecture.txt")

    out, _ = ai.generate_content(
        "s-105",
        "Binary search",
        ["summary", "quiz"],
        "en",
        document_ids=[doc_id],
    )
    assert out["summary"]["type"] == "text"
    assert out["summary"]["source_type"] == "student_upload"
    assert any("sample_lecture.txt" in s for s in out["summary"]["sources"])


def test_selected_material_prevents_external_discovery(tmp_path, monkeypatch):
    """6. Selected material mode prevents invoking external discovery."""
    ai = build_system(str(tmp_path), test_mode=True)
    ai.create_or_load_student("s-106", "Learner", "Computer Science", "en")

    lecture = tmp_path / "sample_lecture.txt"
    lecture.write_text(
        "Binary search requires sorted indexed array. Worst-case is O(log n).",
        encoding="utf-8",
    )
    doc_id = ai.upload_document("s-106", lecture, "sample_lecture.txt")

    called_external = []
    def fake_external(*args, **kwargs):
        called_external.append(True)
        return ([], False, "called", [])

    monkeypatch.setattr(ai, "_external_evidence", fake_external)

    out, _ = ai.generate_content(
        "s-106",
        "Binary search",
        ["summary"],
        "en",
        document_ids=[doc_id],
    )
    assert len(called_external) == 0, "External search must NOT be called when document_ids are provided"
    assert out["summary"]["type"] == "text"


def test_no_material_uses_trusted_external_evidence(tmp_path):
    """7. When no material is selected, trusted external evidence is used."""
    ai = build_system(str(tmp_path), test_mode=True)
    ai.create_or_load_student("s-107", "Learner", "Computer Science", "en")

    out, _ = ai.generate_content(
        "s-107",
        "Binary search",
        ["summary"],
        "en",
        document_ids=None,
    )
    assert out["summary"]["type"] == "text"
    assert out["summary"]["source_type"] == "trusted_external"


def test_no_evidence_safe_refusal(tmp_path):
    """8. Safe refusal when evidence is empty or insufficient."""
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    res = gen.generate("summary", "Quantum Teleportation", "en", evidence=[])
    assert res["type"] == "unavailable"
    assert "Sufficient grounded evidence is required" in res["message"]


def test_no_model_memory_fallback(tmp_path):
    """9. Generator does not fallback to model-memory when evidence is absent."""
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    for kind in ["study_guide", "flashcards", "quiz", "practice", "code"]:
        res = gen.generate(kind, "Unseen Topic", "en", evidence=[])
        assert res["type"] == "unavailable"


def test_summary_contains_no_unsupported_facts(tmp_path):
    """10. Summary synthesizes evidence without introducing unsupported claims."""
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    ev = [
        Evidence(
            chunk_id="c1",
            authority="university",
            text="Binary search requires an array sorted in ascending order. It checks the middle element.",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.2,
        )
    ]
    res = gen.generate("summary", "Binary search", "en", evidence=ev)
    assert res["type"] == "text"
    assert "Binary search" in res["content"]
    assert "sorted" in res["content"].lower()


def test_flashcards_have_valid_front_back_pairs(tmp_path):
    """11. Flashcards contain clear front and concise back for all cards."""
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    ev = [
        Evidence(
            chunk_id="c2",
            authority="university",
            text="Binary search divides array by half. Worst-case runtime is O(log n). Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.2,
        )
    ]
    res = gen.generate("flashcards", "Binary search", "en", evidence=ev)
    assert res["type"] == "text"
    content = res["content"]
    assert "**Card 1**" in content
    assert "**Card 6**" in content
    assert "**Front**:" in content and "**Back**:" in content


def test_duplicate_flashcards_rejected_or_deduplicated():
    """12. Deduplication utility eliminates duplicate flashcard prompts."""
    raw = (
        "Q: What is binary search?\nA: A logarithmic search algorithm.\n\n"
        "Q: What is binary search?\nA: A search algorithm that cuts space in half.\n\n"
        "Q: What is its time complexity?\nA: O(log n)."
    )
    cleaned = ContentGenerator._deduplicate_flashcards(raw)
    assert cleaned.count("What is binary search?") == 1
    assert "What is its time complexity?" in cleaned


def test_practice_expected_answers_grounded(tmp_path):
    """13. Practice questions include verified hints and solutions grounded in evidence."""
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    ev = [
        Evidence(
            chunk_id="c3",
            authority="university",
            text="Binary search checks the middle element of a sorted array and halves the search space.",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.2,
        )
    ]
    res = gen.generate("practice", "Binary search", "en", evidence=ev)
    assert res["type"] == "text"
    assert "Problem 1" in res["content"]
    assert "*Guided Hint*:" in res["content"]
    assert "*Step-by-step Verified Solution*:" in res["content"]


def test_malformed_structured_output_handled_safely(tmp_path):
    """14. Malformed JSON / unexpected exceptions fail safely without crashing."""
    ai = build_system(str(tmp_path), test_mode=True)
    ai.create_or_load_student("s-114", "Learner", "Computer Science", "en")

    # Patch content generator to simulate internal unexpected error on one kind
    def buggy_generate(*args, **kwargs):
        raise RuntimeError("Transient generation fault")

    ai.content.generate = buggy_generate

    out, _ = ai.generate_content("s-114", "Binary search", ["notes"], "en")
    assert "notes" in out
    assert out["notes"]["type"] == "unavailable"
    assert "temporarily unavailable" in out["notes"]["message"]


def test_generation_does_not_mutate_concept_mastery(tmp_path):
    """15. Generating learning resources does NOT change concept_mastery."""
    ai = build_system(str(tmp_path), test_mode=True)
    p = ai.create_or_load_student("s-115", "Learner", "Computer Science", "en")
    p.concept_mastery = {"Binary search": 0.5, "space complexity": 0.2}
    ai.store.save(p)

    ai.generate_content("s-115", "Binary search", ["summary", "study_guide"], "en")
    refreshed = ai.get_student_profile("s-115")
    assert refreshed.concept_mastery == {"Binary search": 0.5, "space complexity": 0.2}


def test_generation_does_not_mutate_unknown_concepts(tmp_path):
    """16. Generating learning resources does NOT change unknown_concepts."""
    ai = build_system(str(tmp_path), test_mode=True)
    p = ai.create_or_load_student("s-116", "Learner", "Computer Science", "en")
    p.unknown_concepts = ["worst-case complexity"]
    ai.store.save(p)

    ai.generate_content("s-116", "worst-case complexity", ["explanation", "quiz"], "en")
    refreshed = ai.get_student_profile("s-116")
    assert refreshed.unknown_concepts == ["worst-case complexity"]


def test_generation_does_not_mutate_assessment_history(tmp_path):
    """17. Generating learning resources does NOT add or mutate assessment_history."""
    ai = build_system(str(tmp_path), test_mode=True)
    p = ai.create_or_load_student("s-117", "Learner", "Computer Science", "en")
    p.assessment_history = [{"assessment_id": "a-1", "score": 0.5}]
    ai.store.save(p)

    ai.generate_content("s-117", "Binary search", ["flashcards"], "en")
    refreshed = ai.get_student_profile("s-117")
    assert len(refreshed.assessment_history) == 1
    assert refreshed.assessment_history[0]["score"] == 0.5


def test_provenance_preserved(tmp_path):
    """18. Evidence sources and source type are preserved in generated resources."""
    ai = build_system(str(tmp_path), test_mode=True)
    ai.create_or_load_student("s-118", "Learner", "Computer Science", "en")

    out, _ = ai.generate_content("s-118", "Binary search", ["notes"], "en")
    notes = out["notes"]
    assert "sources" in notes
    assert "source_type" in notes
    assert notes["source_type"] == "trusted_external"


def test_semantic_differences_across_all_resource_types(tmp_path):
    """19. Verify distinct structures and semantic characteristics across all supported resource types."""
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    ev = [
        Evidence(
            chunk_id="c19",
            authority="university",
            text="Binary search requires an array sorted in ascending order. It evaluates the median element. Worst-case runtime is O(log n). Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.2,
        )
    ]
    kinds = [
        "summary",
        "notes",
        "study_guide",
        "flashcards",
        "quiz",
        "practice",
        "exam",
        "analogy",
        "comparison",
        "question_bank",
        "code",
        "coding_exercise",
        "diagram",
    ]
    results = {}
    for k in kinds:
        res = gen.generate(k, "Binary search", "en", evidence=ev, source_type="upload")
        results[k] = res

    # Check each type has appropriate structure
    assert results["flashcards"]["type"] == "text"
    assert "**Card 1**" in results["flashcards"]["content"]

    assert results["quiz"]["type"] == "text"
    assert "### Answer Key" in results["quiz"]["content"]

    assert results["practice"]["type"] == "text"
    assert "*Step-by-step Verified Solution*:" in results["practice"]["content"]

    assert results["study_guide"]["type"] == "text"
    assert "## 1. Learning Objectives" in results["study_guide"]["content"]

    assert results["exam"]["type"] == "text"
    assert "## Section A: Multiple Choice" in results["exam"]["content"]

    assert results["comparison"]["type"] == "text"
    assert "| Operational Dimension |" in results["comparison"]["content"] or "| Evaluation Criteria |" in results["comparison"]["content"]

    assert results["analogy"]["type"] == "text"
    assert "# Grounded Analogy:" in results["analogy"]["content"]

    assert results["question_bank"]["type"] == "text"
    assert "[Remember]" in results["question_bank"]["content"]

    assert results["code"]["type"] == "text"
    assert "```python" in results["code"]["content"]

    assert results["coding_exercise"]["type"] == "text"
    assert "## Starter Skeleton" in results["coding_exercise"]["content"]

    assert results["diagram"]["type"] == "file"
    assert results["diagram"]["mime"] == "image/svg+xml"
    assert Path(results["diagram"]["path"]).exists()


def test_claim_boundary_hardware_latency_detector_and_sanitizer(tmp_path):
    """20. Detector catches processor speed/clock latency claims; sanitizer strips them while keeping supported facts."""
    ev = [
        Evidence(
            chunk_id="c20",
            authority="university",
            text="Binary search requires a sorted array. It compares the median element. Worst-case is O(log n). Space complexity is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.2,
        )
    ]
    raw_text = (
        "# Deep Dive: binary search\n\n"
        "Binary search operates on a sorted array and achieves O(log n) worst-case time.\n"
        "Importantly, this complexity metric is independent of processor speed or hardware execution latency.\n\n"
        "Space complexity is O(1) auxiliary memory."
    )
    unsupported = ContentGenerator._find_unsupported_claims(raw_text, ev, strict_upload=True)
    assert len(unsupported) > 0
    assert any("processor" in u for u in unsupported)

    sanitized = ContentGenerator._sanitize_unsupported_claims(raw_text, ev, strict_upload=True)
    assert "processor speed" not in sanitized.lower()
    assert "hardware execution latency" not in sanitized.lower()
    assert "O(log n)" in sanitized
    assert "O(1)" in sanitized
    assert ContentGenerator._find_unsupported_claims(sanitized, ev, strict_upload=True) == []


def test_claim_boundary_external_analogy_and_recurrence_sanitization(tmp_path):
    """21. Detector and sanitizer catch ungrounded phone book analogy, Master theorem, and recurrence relations."""
    ev = [
        Evidence(
            chunk_id="c21",
            authority="university",
            text="Binary search operates on a sorted array. It compares the median element. Worst-case is O(log n).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.2,
        )
    ]
    raw_text = (
        "# Analysis: binary search\n\n"
        "Think of binary search like finding a name in a printed telephone directory.\n\n"
        "The recurrence relation T(n) = T(n/2) + O(1) is solved using the Master Theorem.\n\n"
        "Binary search requires an array sorted in ascending order."
    )
    unsupported = ContentGenerator._find_unsupported_claims(raw_text, ev, strict_upload=True)
    assert any("telephone" in u for u in unsupported)
    assert any("Master Theorem" in u or "recurrence" in u for u in unsupported)

    sanitized = ContentGenerator._sanitize_unsupported_claims(raw_text, ev, strict_upload=True)
    assert "telephone directory" not in sanitized.lower()
    assert "master theorem" not in sanitized.lower()
    assert "T(n) = T(n/2)" not in sanitized
    assert "sorted in ascending order" in sanitized
    assert ContentGenerator._find_unsupported_claims(sanitized, ev, strict_upload=True) == []


def test_strict_upload_mode_mock_fallback_all_kinds(tmp_path):
    """22. In strict upload mode, fallback generator produces zero external analogies or ungrounded algorithms."""
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    ev = [
        Evidence(
            chunk_id="c22",
            authority="university",
            text="Binary search requires an array sorted in ascending order. It checks the middle element. Worst-case runtime is O(log n). Iterative space is O(1).",
            source="sample_manual_lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.2,
        )
    ]
    for kind in ["analogy", "comparison", "exam", "question_bank", "summary", "notes"]:
        res = gen.generate(kind, "Binary search", "en", evidence=ev, source_type="upload")
        assert res["type"] == "text"
        content = res["content"]
        # Verify no external ungrounded concepts
        assert "telephone directory" not in content.lower()
        assert "processor speed" not in content.lower()
        assert "master theorem" not in content.lower()
        assert "T(n) = T(n/2)" not in content
        # Verify clean claims
        assert ContentGenerator._find_unsupported_claims(content, ev, strict_upload=True) == []


def test_code_and_coding_exercise_syntax_validity(tmp_path):
    """23. Python code and coding exercise solution compile cleanly as valid Python."""
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    ev = [
        Evidence(
            chunk_id="c23",
            authority="university",
            text="Binary search requires a sorted array. Median is compared. Worst-case is O(log n). Space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.2,
        )
    ]
    code_res = gen.generate("code", "Binary search", "en", evidence=ev, source_type="upload")
    assert code_res["type"] == "text"
    extracted_code = ContentGenerator._extract_python_code(code_res["content"])
    assert extracted_code is not None
    compile(extracted_code, "<test_code>", "exec")

    exercise_res = gen.generate("coding_exercise", "Binary search", "en", evidence=ev, source_type="upload")
    assert exercise_res["type"] == "text"
    extracted_exercise = ContentGenerator._extract_python_code(exercise_res["content"])
    assert extracted_exercise is not None
    compile(extracted_exercise, "<test_exercise>", "exec")


def test_practice_and_quiz_answer_keys_are_grounded(tmp_path):
    """24. Practice problems and quiz answer keys contain verified solutions directly grounded in evidence."""
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    ev = [
        Evidence(
            chunk_id="c24",
            authority="university",
            text="Binary search requires a sorted indexed array. Halves search interval at each step. Worst-case is O(log n). Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.2,
        )
    ]
    practice = gen.generate("practice", "Binary search", "en", evidence=ev, source_type="upload")
    quiz = gen.generate("quiz", "Binary search", "en", evidence=ev, source_type="upload")

    assert "O(log n)" in practice["content"]
    assert "Step-by-step Verified Solution" in practice["content"]

    assert "### Answer Key" in quiz["content"]
    assert "O(log n)" in quiz["content"]
    assert "O(1)" in quiz["content"]


def test_generic_claim_validator_exact_fixture_rejects_unsupported_definition(tmp_path):
    """25. Exact fixture test: reject general theoretical definition of worst-case complexity."""
    ev = [
        Evidence(
            chunk_id="c_manual",
            authority="university",
            text=(
                "Course: CS101 Introduction to Algorithms & Systems\n"
                "Topic: Divide-and-Conquer Search and Computational Limits\n\n"
                "Core Theoretical Foundations:\n"
                "1. Binary Search is a fundamental search algorithm operating on a sorted indexed array.\n"
                "2. In each iteration, the algorithm compares the target key against the median element of the current subarray.\n"
                "3. If the target equals the median element, its index is returned immediately.\n"
                "4. If the target is smaller, the search domain is restricted to the lower half; if larger, to the upper half.\n"
                "5. The worst-case time complexity is O(log n), and space complexity is O(1) for iterative implementations.\n"
                "6. The indispensable precondition before binary search can be applied is that the data sequence must be strictly sorted.\n\n"
                "Synthetic Manual Verification Marker:\n"
                "UniqueManualProofToken#9824: In the Edunation experimental lab evaluation, the cryogenic array benchmark achieved an exact latency of 4.81 microseconds across ten million simulated lookups."
            ),
            source="sample_manual_lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    unsupported_statement = (
        "Worst-case complexity is a theoretical measure used to describe the maximum "
        "amount of resources an algorithm can take as a function of input size."
    )
    classification = ContentGenerator._classify_claim(unsupported_statement, ev)
    assert classification == "UNSUPPORTED", f"Expected UNSUPPORTED, got {classification}"

    unsupported_list = ContentGenerator._find_unsupported_claims(unsupported_statement, ev)
    assert len(unsupported_list) > 0


def test_generic_claim_validator_adversarial_claims_rejected(tmp_path):
    """26. Adversarial tests: reject claims not in any blacklist without keyword matching."""
    ev = [
        Evidence(
            chunk_id="c_manual_adv",
            authority="university",
            text=(
                "Binary Search operates on a sorted indexed array by comparing the target with the median element. "
                "Worst-case time complexity is O(log n), and iterative space complexity is O(1)."
            ),
            source="sample_manual_lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    adversarial_claims = [
        "Worst-case complexity measures maximum resource usage as input grows.",
        "Big-O abstracts away hardware speed.",
        "Binary Search is optimal among all comparison-based search algorithms.",
        "Binary Search was invented in the 20th century.",
        "Logarithmic algorithms scale efficiently to billions of items.",
    ]
    for claim in adversarial_claims:
        classification = ContentGenerator._classify_claim(claim, ev)
        assert classification == "UNSUPPORTED", f"Adversarial claim failed rejection: {claim} -> {classification}"


def test_generic_claim_validator_valid_paraphrases_accepted(tmp_path):
    """27. Directly implied reasoning and valid paraphrases are accepted."""
    ev = [
        Evidence(
            chunk_id="c_manual_val",
            authority="university",
            text=(
                "Binary Search is a fundamental search algorithm operating on a sorted indexed array. "
                "In each iteration, the algorithm compares the target key against the median element of the current subarray. "
                "If the target equals the median element, its index is returned immediately. "
                "If target is smaller, the search domain is restricted to the lower half; if larger, to the upper half. "
                "The worst-case time complexity is O(log n), and space complexity is O(1) for iterative implementations. "
                "The precondition is that the data sequence must be strictly sorted."
            ),
            source="sample_manual_lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    valid_paraphrases = [
        "Binary Search operates on a sorted indexed array.",
        "Binary search requires an array sorted in ascending order.",
        "Binary Search narrows the search domain during each comparison.",
        "Binary Search has logarithmic worst-case time complexity.",
        "Space complexity is O(1) auxiliary memory.",
        "The data sequence must be strictly sorted before binary search can be applied.",
        "If the target matches the median, it returns immediately.",
    ]
    for paraphrase in valid_paraphrases:
        classification = ContentGenerator._classify_claim(paraphrase, ev)
        assert classification in ("SUPPORTED", "DIRECTLY_IMPLIED"), f"Valid paraphrase rejected: {paraphrase} -> {classification}"


def test_generic_claim_validator_pedagogical_framing_not_rejected(tmp_path):
    """28. Pedagogical text, section headings, and transitions are not rejected as factual claims."""
    ev = [
        Evidence(
            chunk_id="c_ped",
            authority="university",
            text="Binary search operates on a sorted array. Worst-case is O(log n).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    pedagogical_examples = [
        "### 1. Conceptual Framework & Theoretical Principles",
        "Building upon your mastery of the fundamentals, we will analyze the deeper mechanics of Binary Search.",
        "Let us examine this step by step.",
        "Consider the following concrete input:",
        "*Pedagogical note*: Foundational learning mode active — no prior mastery assumed.",
        "## Problem 1 (Foundational — Core Principle)",
        "*Guided Hint*: Review the basic definitions in the evidence.",
        "**Card 1**",
        "**Front**: What is the primary operational requirement of Binary search?",
    ]
    for ped in pedagogical_examples:
        classification = ContentGenerator._classify_claim(ped, ev)
        assert classification in ("SUPPORTED", "DIRECTLY_IMPLIED"), f"Pedagogical text incorrectly flagged: {ped} -> {classification}"


def test_generic_claim_validator_qa_resources_both_premise_and_answer(tmp_path):
    """29. Question and answer validation: both premise and answer key must be grounded."""
    ev = [
        Evidence(
            chunk_id="c_qa",
            authority="university",
            text="Binary search requires a sorted indexed array. Median element is compared. Worst-case is O(log n). Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    # Grounded Q/A pair
    grounded_qa = (
        "1. [Remember] State the worst-case time complexity of Binary search.\n"
        "   *Answer*: O(log n).\n"
    )
    unsupported_grounded = ContentGenerator._find_unsupported_claims(grounded_qa, ev)
    assert len(unsupported_grounded) == 0

    # Ungrounded premise: Sorting takes O(n log n)
    ungrounded_premise = (
        "5. [Evaluate] Assess the trade-off of sorting an unsorted array where sorting takes O(n log n).\n"
        "   *Answer*: Sorting takes O(n log n), which is slower than a single linear scan O(n).\n"
    )
    unsupported_ungrounded = ContentGenerator._find_unsupported_claims(ungrounded_premise, ev)
    assert len(unsupported_ungrounded) > 0


def test_generic_claim_validator_end_to_end_sanitization_removes_unsupported(tmp_path):
    """30. End-to-end sanitization cleanly strips ungrounded claims and empty sections while preserving grounded content."""
    ev = [
        Evidence(
            chunk_id="c_e2e",
            authority="university",
            text=(
                "Binary Search is a fundamental search algorithm operating on a sorted indexed array. "
                "In each iteration, the algorithm compares the target key against the median element of the current subarray. "
                "If the target equals the median element, its index is returned immediately. "
                "If the target is smaller, the search domain is restricted to the lower half; if larger, to the upper half. "
                "The worst-case time complexity is O(log n), and space complexity is O(1) for iterative implementations."
            ),
            source="sample_manual_lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    mixed_content = (
        "### 1. Conceptual Framework & Theoretical Principles\n"
        "Worst-case complexity is a theoretical measure used to describe the maximum amount of resources an algorithm can take as a function of input size.\n\n"
        "### 2. Operational Invariants in Binary Search\n"
        "Binary Search is a fundamental search algorithm operating on a sorted indexed array.\n"
        "In each iteration, the algorithm compares the target key against the median element of the current subarray.\n"
        "Big-O abstracts away hardware speed.\n"
        "If the target equals the median element, its index is returned immediately.\n\n"
        "### 3. Complexity Boundaries\n"
        "The worst-case time complexity is O(log n), and space complexity is O(1) for iterative implementations.\n"
    )
    sanitized = ContentGenerator._sanitize_unsupported_claims(mixed_content, ev)

    assert "theoretical measure" not in sanitized
    assert "hardware speed" not in sanitized
    assert "Binary Search is a fundamental search algorithm" in sanitized
    assert "O(log n)" in sanitized
    assert "O(1)" in sanitized

    # Revalidation proves zero remaining unsupported claims
    remaining = ContentGenerator._find_unsupported_claims(sanitized, ev)
    assert len(remaining) == 0


def test_contract_sanitizer_cannot_leave_problem_without_question(tmp_path):
    """31. Sanitizer cannot leave a Problem without an explicit Question."""
    ev = [
        Evidence(
            chunk_id="c_sc",
            authority="university",
            text="Binary Search iterative space is O(1). Array must be strictly sorted.",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    broken_content = (
        "# Practice Problem Set: Binary Search\n\n"
        "## Problem 1 (Foundational)\n"
        "*Guided Hint*: Check point 5 in the evidence.\n"
        "*Step-by-step Verified Solution*: Iterative space is O(1).\n\n"
        "## Problem 2 (Applied)\n"
        "What is the iterative space complexity of Binary Search?\n"
        "*Guided Hint*: Check point 5 in the evidence.\n"
        "*Step-by-step Verified Solution*: Iterative space is O(1).\n"
    )
    is_valid, normalized, errors = ContentGenerator._validate_resource_contract(
        "practice", broken_content, ev, "Binary Search"
    )
    assert "Problem 1 is missing an explicit Question" in str(errors)
    assert "Problem 1 (Foundational)" not in normalized
    assert "## Problem 1 (Applied)" in normalized
    assert "What is the iterative space complexity of Binary Search?" in normalized


def test_contract_sanitizer_cannot_leave_problem_without_hint(tmp_path):
    """32. Sanitizer cannot leave a Problem without a Guided Hint."""
    ev = [
        Evidence(
            chunk_id="c_sc",
            authority="university",
            text="Binary Search iterative space is O(1). Array must be strictly sorted.",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    broken_content = (
        "# Practice Problem Set: Binary Search\n\n"
        "## Problem 1 (Foundational)\n"
        "What precondition must hold before Binary Search can be applied?\n"
        "*Step-by-step Verified Solution*: Array must be strictly sorted.\n\n"
        "## Problem 2 (Applied)\n"
        "What is the iterative space complexity of Binary Search?\n"
        "*Guided Hint*: Check point 5 in the evidence.\n"
        "*Step-by-step Verified Solution*: Iterative space is O(1).\n"
    )
    is_valid, normalized, errors = ContentGenerator._validate_resource_contract(
        "practice", broken_content, ev, "Binary Search"
    )
    assert "Problem 1 is missing a Guided Hint" in str(errors)
    assert "## Problem 1 (Applied)" in normalized


def test_contract_sanitizer_cannot_leave_problem_without_solution(tmp_path):
    """33. Sanitizer cannot leave a Problem without a Solution."""
    ev = [
        Evidence(
            chunk_id="c_sc",
            authority="university",
            text="Binary Search iterative space is O(1). Array must be strictly sorted.",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    broken_content = (
        "# Practice Problem Set: Binary Search\n\n"
        "## Problem 1 (Foundational)\n"
        "What precondition must hold before Binary Search can be applied?\n"
        "*Guided Hint*: Review the array structure requirement.\n\n"
        "## Problem 2 (Applied)\n"
        "What is the iterative space complexity of Binary Search?\n"
        "*Guided Hint*: Check point 5 in the evidence.\n"
        "*Step-by-step Verified Solution*: Iterative space is O(1).\n"
    )
    is_valid, normalized, errors = ContentGenerator._validate_resource_contract(
        "practice", broken_content, ev, "Binary Search"
    )
    assert "Problem 1 is missing a Step-by-step Solution" in str(errors)
    assert "## Problem 1 (Applied)" in normalized


def test_contract_unsupported_problem_item_dropped_as_a_whole(tmp_path):
    """34. Unsupported problem item is dropped as a whole unit, not left half-stripped."""
    ev = [
        Evidence(
            chunk_id="c_sc",
            authority="university",
            text="Binary Search operates on a sorted array. Iterative space complexity is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    content = (
        "# Practice Problem Set: space complexity\n\n"
        "## Problem 1 (Theoretical Definition)\n"
        "Space complexity quantifies the total auxiliary heap and stack frames required by dynamic memory.\n"
        "*Guided Hint*: Review memory allocation in computer systems.\n"
        "*Step-by-step Verified Solution*: Dynamic memory allocates pointers on heap buffers.\n\n"
        "## Problem 2 (Grounded Invariant)\n"
        "What is the space complexity for iterative implementations of Binary Search?\n"
        "*Guided Hint*: Check the complexity bounds in the evidence.\n"
        "*Step-by-step Verified Solution*: Iterative space complexity is O(1).\n"
    )
    is_valid, normalized, errors = ContentGenerator._validate_resource_contract(
        "practice", content, ev, "space complexity"
    )
    assert not is_valid
    assert "Problem 1 Question contains unsupported claims" in str(errors)
    assert "Problem 1 (Theoretical Definition)" not in normalized
    assert "heap and stack" not in normalized
    assert "## Problem 1 (Grounded Invariant)" in normalized
    assert "Iterative space complexity is O(1)" in normalized


def test_contract_fewer_complete_problems_allowed_when_evidence_limited(tmp_path):
    """35. Fewer complete problems returned when evidence safely supports limited items."""
    ev = [
        Evidence(
            chunk_id="c_sc",
            authority="university",
            text="Binary Search iterative space is O(1). Array must be strictly sorted.",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    res = gen.generate("practice", "space complexity", "en", evidence=ev, source_type="upload")
    assert res["type"] == "text"
    content = res["content"]
    assert "Problem 1" in content
    assert "*Guided Hint*:" in content
    assert "*Step-by-step Verified Solution*:" in content
    unsupported = ContentGenerator._find_unsupported_claims(content, ev)
    assert len(unsupported) == 0


def test_contract_no_content_padding_with_unrelated_grounded_evidence(tmp_path):
    """36. No content padding with unrelated grounded facts."""
    ev = [
        Evidence(
            chunk_id="c_manual",
            authority="university",
            text=open("data/manual_test/sample_manual_lecture.txt").read(),
            source="sample_manual_lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    res = gen.generate("practice", "space complexity", "en", evidence=ev, source_type="upload")
    content = res["content"]
    assert "latency" not in content.lower()
    assert "cryogenic" not in content.lower()
    assert "simulated lookups" not in content.lower()


def test_contract_latency_marker_not_injected_into_unrelated_space_complexity(tmp_path):
    """37. Latency benchmark marker is not injected into unrelated space-complexity problems."""
    ev = [
        Evidence(
            chunk_id="c_manual",
            authority="university",
            text=open("data/manual_test/sample_manual_lecture.txt").read(),
            source="sample_manual_lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    unrelated_problem = (
        "## Problem 2 (Benchmark Evaluation)\n"
        "In the Edunation experimental lab evaluation, the cryogenic array benchmark achieved an exact latency of 4.81 microseconds across ten million simulated lookups. How does this impact cache memory?\n"
        "*Guided Hint*: Review the synthetic token in the notes.\n"
        "*Step-by-step Verified Solution*: The latency benchmark measures memory lookup time.\n"
    )
    is_valid, normalized, errors = ContentGenerator._validate_resource_contract(
        "practice", unrelated_problem, ev, "space complexity"
    )
    assert not is_valid
    assert any("irrelevant" in err.lower() for err in errors)


def test_contract_flashcard_missing_front_back_rejected(tmp_path):
    """38. Flashcard missing Front or Back is rejected as a whole unit."""
    ev = [
        Evidence(
            chunk_id="c_fc",
            authority="university",
            text="Binary search requires a sorted indexed array. Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    cards_text = (
        "# Flashcards: Binary search\n\n"
        "**Card 1**\n"
        "**Front**: What is the space complexity of iterative binary search?\n"
        "**Back**: Iterative space is O(1).\n\n"
        "**Card 2**\n"
        "**Front**: What is the precondition before binary search can run?\n\n"
        "**Card 3**\n"
        "**Back**: Sorted indexed array.\n"
    )
    is_valid, normalized, errors = ContentGenerator._validate_resource_contract(
        "flashcards", cards_text, ev, "Binary search"
    )
    assert not is_valid
    assert "**Card 1**" in normalized
    assert "Iterative space is O(1)" in normalized
    assert "**Card 2**" not in normalized


def test_contract_quiz_missing_answer_rejected(tmp_path):
    """39. Quiz question without an answer in Answer Key is rejected."""
    ev = [
        Evidence(
            chunk_id="c_q",
            authority="university",
            text="Binary search operates on sorted array. Worst-case is O(log n). Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    quiz_text = (
        "# Practice Quiz: Binary search\n\n"
        "1. What is the worst-case runtime of Binary search?\n"
        "2. What is the space complexity of iterative Binary search?\n"
        "3. What is the external hardware architecture?\n\n"
        "### Answer Key\n"
        "1. Worst-case is O(log n).\n"
        "2. Iterative space is O(1).\n"
    )
    is_valid, normalized, errors = ContentGenerator._validate_resource_contract(
        "quiz", quiz_text, ev, "Binary search"
    )
    assert not is_valid
    assert "Question 3 has no matching answer" in str(errors)
    assert "1. What is the worst-case runtime" in normalized
    assert "2. What is the space complexity" in normalized
    assert "hardware architecture" not in normalized


def test_contract_exam_incomplete_section_handled_safely(tmp_path):
    """40. Exam missing Answer Key or required sections is flagged and normalized safely."""
    ev = [
        Evidence(
            chunk_id="c_ex",
            authority="university",
            text="Binary search requires sorted array. Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    broken_exam = (
        "# Formal Examination: Binary search\n\n"
        "## Section A: Multiple Choice\n"
        "1. What is iterative space complexity?\n"
        "A) O(1) B) O(n) C) O(n^2) D) O(log n)\n"
    )
    is_valid, normalized, errors = ContentGenerator._validate_resource_contract(
        "exam", broken_exam, ev, "Binary search"
    )
    assert not is_valid
    assert any("answer key" in err.lower() for err in errors)


def test_contract_question_bank_qa_remain_paired(tmp_path):
    """41. Question bank items missing answer or question are dropped, keeping valid pairs intact."""
    ev = [
        Evidence(
            chunk_id="c_qb",
            authority="university",
            text="Binary search requires sorted array. Worst-case is O(log n). Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    bank_text = (
        "# Question Bank: Binary search\n\n"
        "1. [Remember] State the worst-case runtime of Binary search.\n"
        "   *Answer*: Worst-case is O(log n).\n"
        "2. [Understand] Explain hardware CPU pipeline register hazards.\n"
        "3. [Apply] What is the space complexity for iterative Binary search?\n"
        "   *Answer*: Iterative space is O(1).\n"
    )
    is_valid, normalized, errors = ContentGenerator._validate_resource_contract(
        "question_bank", bank_text, ev, "Binary search"
    )
    assert not is_valid
    assert "1. [Remember]" in normalized
    assert "2. [Apply]" in normalized
    assert "hardware CPU pipeline" not in normalized


def test_contract_revalidated_after_grounding_sanitization(tmp_path):
    """42. Contract validation runs after grounding sanitization, preventing orphan headings or half-items."""
    ev = [
        Evidence(
            chunk_id="c_val",
            authority="university",
            text="Binary search operates on sorted array. Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    content_with_unsupported_question = (
        "# Practice Problem Set: Binary search\n\n"
        "## Problem 1 (Advanced Hardware Design)\n"
        "Describe how speculative execution hazards interact with branch target buffers in modern superscalar processors.\n"
        "*Guided Hint*: Review CPU branch prediction mechanisms.\n"
        "*Step-by-step Verified Solution*: Branch predictors evaluate speculative execution paths.\n\n"
        "## Problem 2 (Iterative Space)\n"
        "What is the space complexity for iterative implementations of Binary search?\n"
        "*Guided Hint*: Check point 5 in the evidence.\n"
        "*Step-by-step Verified Solution*: Iterative space is O(1).\n"
    )
    is_valid, normalized, errors = ContentGenerator._validate_resource_contract(
        "practice", content_with_unsupported_question, ev, "Binary search"
    )
    assert not is_valid
    assert "## Problem 1 (Advanced Hardware Design)" not in normalized
    assert "speculative execution" not in normalized
    assert "## Problem 1 (Iterative Space)" in normalized


def test_grounding_revalidated_after_structured_repair(tmp_path):
    """43. Grounding is strictly re-evaluated after structured repair attempts."""
    ev = [
        Evidence(
            chunk_id="c_rep",
            authority="university",
            text="Binary search operates on sorted array. Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    repaired_with_hallucination = (
        "# Practice Problem Set: Binary search\n\n"
        "## Problem 1 (Space Complexity)\n"
        "What is the space complexity for iterative Binary search?\n"
        "*Guided Hint*: Check the evidence.\n"
        "*Step-by-step Verified Solution*: In 1960, John von Neumann proved iterative space is O(1).\n"
    )
    is_valid, normalized, errors = ContentGenerator._validate_resource_contract(
        "practice", repaired_with_hallucination, ev, "Binary search"
    )
    assert not is_valid
    assert any("unsupported" in err.lower() for err in errors)


def test_bounded_repair_maximum_one_retry(tmp_path):
    """44. Generator bounded repair executes at most one retry, never looping infinitely."""
    ev = [
        Evidence(
            chunk_id="c_rep",
            authority="university",
            text="Binary search operates on sorted array. Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    call_count = 0
    class AlwaysHallucinatingLLM:
        def generate_prompt(self, system_prompt, user_prompt):
            nonlocal call_count
            call_count += 1
            return (
                "# Practice Problem Set: Binary search\n\n"
                "## Problem 1 (Broken)\n"
                "*Guided Hint*: Only hint, no question.\n"
                "*Step-by-step Verified Solution*: Hallucinated external fact 1970.\n"
            )

    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    mock_llm = AlwaysHallucinatingLLM()
    res = gen.generate("practice", "space complexity", "en", llm=mock_llm, evidence=ev, source_type="upload")

    # Initial call (1) + at most ONE repair retry (1) = 2 total calls
    assert call_count <= 2
    assert res["type"] == "text"
    assert "Problem 1" in res["content"]
    assert "*Guided Hint*:" in res["content"]
    assert "*Step-by-step Verified Solution*:" in res["content"]


def test_final_artifact_must_satisfy_grounding_and_contract(tmp_path):
    """45. Final delivered artifact across all 14 types satisfies BOTH grounding and contract."""
    ev = [
        Evidence(
            chunk_id="c_all",
            authority="university",
            text=open("data/manual_test/sample_manual_lecture.txt").read(),
            source="sample_manual_lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    kinds = [
        "explanation", "summary", "notes", "study_guide",
        "flashcards", "quiz", "practice", "exam",
        "analogy", "comparison", "question_bank",
        "code", "coding_exercise", "diagram"
    ]
    for k in kinds:
        res = gen.generate(k, "space complexity", "en", evidence=ev, source_type="upload")
        assert res["type"] in ("text", "file"), f"Kind {k} failed type check"
        if res["type"] == "text":
            content = res["content"]
            unsupported = ContentGenerator._find_unsupported_claims(content, ev)
            assert len(unsupported) == 0, f"Kind {k} has unsupported claims: {unsupported}"
            assert "4. Answer:" not in content
            assert "cryogenic" not in content.lower()
            assert "simulated lookups" not in content.lower()


def test_profile_state_remains_unchanged_during_study_tools_generation(tmp_path):
    """46. Student profile state is not mutated by study tool content generation."""
    import json
    profile_path = Path("data/profiles/student-001.json")
    original_data = json.loads(profile_path.read_text(encoding="utf-8"))

    ev = [
        Evidence(
            chunk_id="c_prof",
            authority="university",
            text="Binary search requires a sorted indexed array. Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    gen.generate("practice", "space complexity", "en", evidence=ev, source_type="upload")
    gen.generate("explanation", "space complexity", "en", evidence=ev, source_type="upload")

    current_data = json.loads(profile_path.read_text(encoding="utf-8"))
    assert original_data == current_data


def test_diagram_artifact_contract_and_download_payload(tmp_path):
    """47. Diagram artifact satisfies explicit contract: preview_svg, download_bytes, filename, mime_type."""
    ev = [
        Evidence(
            chunk_id="c_diag",
            authority="university",
            text="Binary search requires a sorted indexed array. It compares the median element and halves search space.",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    res = gen.generate("diagram", "binary search", "en", evidence=ev, source_type="upload")

    assert res["type"] == "file"
    assert res["kind"] == "diagram"
    assert res["preview_svg"] is not None
    assert isinstance(res["preview_svg"], str)
    assert "<svg" in res["preview_svg"]

    assert res["download_bytes"] is not None
    assert isinstance(res["download_bytes"], bytes)
    assert len(res["download_bytes"]) > 0
    assert res["download_bytes"].startswith(b"<svg")

    assert res["filename"].endswith(".svg")
    assert res["mime_type"] == "image/svg+xml"
    assert res["mime"] == "image/svg+xml"
    assert Path(res["path"]).exists()


def test_presentation_artifact_contract_and_download_payload(tmp_path):
    """48. Presentation artifact satisfies explicit contract: valid PPTX bytes, ZIP signature, python-pptx loadable, rewound BytesIO."""
    import io
    from pptx import Presentation

    ev = [
        Evidence(
            chunk_id="c_ppt",
            authority="university",
            text="Binary search algorithm operates on sorted array. O(log n) time and O(1) iterative space.",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    res = gen.generate("presentation", "binary search", "en", evidence=ev, source_type="upload")

    assert res["type"] == "file"
    assert res["kind"] == "presentation"
    assert res["download_bytes"] is not None
    assert isinstance(res["download_bytes"], bytes)
    assert len(res["download_bytes"]) > 0

    # OpenXML presentations are ZIP archives starting with PK\x03\x04
    assert res["download_bytes"][:4] == b"PK\x03\x04"

    # Verify python-pptx can load from the in-memory download_bytes
    prs = Presentation(io.BytesIO(res["download_bytes"]))
    assert len(prs.slides) >= 6

    assert res["filename"].endswith(".pptx")
    assert res["mime_type"] == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    assert res["mime"] == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    assert Path(res["path"]).exists()
    assert Path(res["path"]).stat().st_size == len(res["download_bytes"])


def test_ui_download_button_payload_wiring_contract(tmp_path):
    """49. UI download button adapter receives actual bytes/string without temporary URL dependencies."""
    import io
    from pptx import Presentation

    ev = [
        Evidence(
            chunk_id="c_ui",
            authority="university",
            text="Binary search requires sorted array. Runtime is O(log n).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    diag_res = gen.generate("diagram", "binary search", "en", evidence=ev, source_type="upload")
    ppt_res = gen.generate("presentation", "binary search", "en", evidence=ev, source_type="upload")

    # Simulate UI extraction logic in app/streamlit_app.py
    # Diagram UI contract
    diag_bytes = diag_res.get("download_bytes")
    diag_filename = diag_res.get("filename")
    diag_mime = diag_res.get("mime_type")

    assert isinstance(diag_bytes, bytes)
    assert len(diag_bytes) > 0
    assert diag_filename.endswith(".svg")
    assert diag_mime == "image/svg+xml"
    # No temporary URL or wrapper
    assert not isinstance(diag_bytes, str) or not diag_bytes.startswith("http")

    # Presentation UI contract
    ppt_bytes = ppt_res.get("download_bytes")
    ppt_filename = ppt_res.get("filename")
    ppt_mime = ppt_res.get("mime_type")

    assert isinstance(ppt_bytes, bytes)
    assert len(ppt_bytes) > 0
    assert ppt_bytes.startswith(b"PK\x03\x04")
    assert ppt_filename.endswith(".pptx")
    assert ppt_mime == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    # Verify Presentation can be re-loaded
    prs = Presentation(io.BytesIO(ppt_bytes))
    assert len(prs.slides) >= 6

    # Verify Streamlit UI contract in app/streamlit_app.py
    ui_source = Path("app/streamlit_app.py").read_text(encoding="utf-8")
    assert 'label="Download Diagram"' in ui_source
    assert 'label="Download Presentation Deck"' in ui_source
    assert "data=download_bytes" in ui_source
    assert "image/svg+xml" in ui_source
    assert "application/vnd.openxmlformats-officedocument.presentationml.presentation" in ui_source
    assert "http://" not in ui_source and "localhost:" not in ui_source


def test_failed_binary_search_artifacts_rejected_by_semantic_validators():
    """50. Semantic validators catch the exact failure modes in previous binary_search artifacts."""
    with open("data/manual_test/sample_manual_lecture.txt") as f:
        ev_text = f.read()

    ev = [
        Evidence(
            chunk_id="c_manual",
            authority="course",
            text=ev_text,
            source="sample_manual_lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]

    # 1. Test previous failed PPTX
    failed_pptx_path = Path("data/artifacts/binary_search_6f4be19c.pptx")
    if failed_pptx_path.exists():
        pptx_bytes = failed_pptx_path.read_bytes()
        valid, errors = ContentGenerator._validate_rendered_pptx(pptx_bytes, "binary search", ev)
        assert not valid, "Previous failed binary search PPTX must be rejected by semantic validator"
        assert any("neural marker" in e or "unsupported" in e or "matrix shapes" in e for e in errors)

    # 2. Test previous failed SVG
    failed_svg_path = Path("data/artifacts/binary_search_22ef0385.svg")
    if failed_svg_path.exists():
        svg_text = failed_svg_path.read_text(encoding="utf-8")
        valid, errors = ContentGenerator._validate_rendered_svg(svg_text, "binary search", ev)
        assert not valid, "Previous failed binary search SVG must be rejected by semantic validator"
        assert any("generic placeholder" in e or "unsupported" in e for e in errors)


def test_grounded_binary_search_presentation_generation_and_validation(tmp_path):
    """51. Newly generated binary search PPTX is grounded, valid, and free of neural network contamination."""
    import io
    from pptx import Presentation

    with open("data/manual_test/sample_manual_lecture.txt") as f:
        ev_text = f.read()

    ev = [
        Evidence(
            chunk_id="c_manual",
            authority="course",
            text=ev_text,
            source="sample_manual_lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]

    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    ppt_res = gen.generate("presentation", "binary search", "en", evidence=ev, source_type="upload")

    assert ppt_res["type"] == "file"
    assert ppt_res["kind"] == "presentation"
    pptx_bytes = ppt_res["download_bytes"]
    assert pptx_bytes.startswith(b"PK\x03\x04")

    # Validate with semantic validator
    valid, errors = ContentGenerator._validate_rendered_pptx(pptx_bytes, "binary search", ev)
    assert valid, f"Generated binary search PPTX must pass semantic validator: {errors}"
    assert len(errors) == 0

    # Extract all text and inspect content
    prs = Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) >= 6

    all_text = []
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    all_text.append(p.text.lower())

    full_corpus = " ".join(all_text)
    # Required topic concepts must be present
    assert "binary search" in full_corpus
    assert "sorted" in full_corpus
    assert "median" in full_corpus
    assert "o(log n)" in full_corpus or "log" in full_corpus
    assert "o(1)" in full_corpus

    # Prohibited neural network concepts must NOT be present
    prohibited = [
        "weight vector", "weights", "hidden unit", "hidden representation",
        "matrix shapes & checks", "matrix shapes", "activation function",
        "w_j", "beta", "forward computation"
    ]
    for p in prohibited:
        assert p not in full_corpus, f"Found prohibited domain token '{p}' in binary search presentation!"


def test_grounded_binary_search_diagram_generation_and_validation(tmp_path):
    """52. Newly generated binary search SVG is grounded, valid, and free of generic placeholder nodes."""
    with open("data/manual_test/sample_manual_lecture.txt") as f:
        ev_text = f.read()

    ev = [
        Evidence(
            chunk_id="c_manual",
            authority="course",
            text=ev_text,
            source="sample_manual_lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]

    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    diag_res = gen.generate("diagram", "binary search", "en", evidence=ev, source_type="upload")

    assert diag_res["type"] == "file"
    assert diag_res["kind"] == "diagram"
    svg_str = diag_res["download_bytes"].decode("utf-8")

    # Validate with semantic validator
    valid, errors = ContentGenerator._validate_rendered_svg(svg_str, "binary search", ev)
    assert valid, f"Generated binary search SVG must pass semantic validator: {errors}"
    assert len(errors) == 0

    # Verify no generic placeholder text appears
    placeholders = [
        "core idea", "mechanism", "how the concept transforms or processes information",
        "what the process produces", "connect the result back to the learning goal"
    ]
    for ph in placeholders:
        assert ph not in svg_str.lower(), f"Found generic placeholder '{ph}' in SVG!"

    # Verify topic-specific facts appear
    assert "sorted" in svg_str.lower()
    assert "median" in svg_str.lower()
    assert "half" in svg_str.lower()


def test_diagram_viewbox_integrity_and_node_bounds(tmp_path):
    """53. SVG viewBox strictly bounds all boxes, labels, and text without clipping."""
    import xml.etree.ElementTree as ET

    with open("data/manual_test/sample_manual_lecture.txt") as f:
        ev_text = f.read()

    ev = [
        Evidence(
            chunk_id="c_manual",
            authority="course",
            text=ev_text,
            source="sample_manual_lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]

    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    diag_res = gen.generate("diagram", "binary search", "en", evidence=ev, source_type="upload")
    svg_str = diag_res["download_bytes"].decode("utf-8")

    root = ET.fromstring(svg_str)
    vb = root.attrib["viewBox"]
    vx, vy, vw, vh = map(float, vb.split())

    # Check rect elements
    for rect in root.iter("{http://www.w3.org/2000/svg}rect"):
        x = float(rect.attrib["x"])
        y = float(rect.attrib["y"])
        w = float(rect.attrib["width"])
        h = float(rect.attrib["height"])
        assert x >= vx, f"Rect x {x} < viewBox min_x {vx}"
        assert x + w <= vx + vw + 0.1, f"Rect right {x + w} > viewBox max_x {vx + vw}"
        assert y >= vy, f"Rect y {y} < viewBox min_y {vy}"
        assert y + h <= vy + vh + 0.1, f"Rect bottom {y + h} > viewBox max_y {vy + vh}"

    # Check text elements
    for text in root.iter("{http://www.w3.org/2000/svg}text"):
        x = float(text.attrib["x"])
        y = float(text.attrib["y"])
        assert vx <= x <= vx + vw, f"Text x {x} outside viewBox [{vx}, {vx + vw}]"
        assert vy <= y <= vy + vh, f"Text y {y} outside viewBox [{vy}, {vy + vh}]"


def test_cross_topic_contamination_prevention(tmp_path):
    """54. Consecutive generations across different topics remain strictly isolated without template leakage."""
    import io
    from pptx import Presentation

    neural_ev = [
        Evidence(
            chunk_id="c_nn",
            authority="course",
            text="Forward propagation computes z = φ(W x) and output f = β^T z using learned weight matrices.",
            source="neural.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    algo_ev = [
        Evidence(
            chunk_id="c_algo",
            authority="course",
            text="Binary search requires a strictly sorted indexed array, comparing median element in O(log n) time.",
            source="algorithms.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]

    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))

    # Step 1: Generate for Topic A (neural forward propagation)
    ppt_nn = gen.generate("presentation", "forward propagation", "en", evidence=neural_ev, source_type="upload")
    diag_nn = gen.generate("diagram", "forward propagation", "en", evidence=neural_ev, source_type="upload")

    prs_nn = Presentation(io.BytesIO(ppt_nn["download_bytes"]))
    nn_ppt_text = " ".join(
        p.text.lower()
        for s in prs_nn.slides
        for sh in s.shapes
        if sh.has_text_frame
        for p in sh.text_frame.paragraphs
    )
    assert "weight" in nn_ppt_text

    # Step 2: Immediately generate for Topic B (binary search)
    ppt_bs = gen.generate("presentation", "binary search", "en", evidence=algo_ev, source_type="upload")
    diag_bs = gen.generate("diagram", "binary search", "en", evidence=algo_ev, source_type="upload")

    # Step 3: Assert Topic B contains ZERO Topic A content
    prs_bs = Presentation(io.BytesIO(ppt_bs["download_bytes"]))
    bs_ppt_text = " ".join(
        p.text.lower()
        for s in prs_bs.slides
        for sh in s.shapes
        if sh.has_text_frame
        for p in sh.text_frame.paragraphs
    )
    import xml.etree.ElementTree as ET
    root_bs = ET.fromstring(diag_bs["download_bytes"].decode("utf-8"))
    bs_svg_text = " ".join(
        elem.text.lower()
        for elem in root_bs.iter()
        if elem.tag.endswith("text") and elem.text
    )

    neural_tokens = ["weight", "hidden", "activation", "top-layer", "beta", "φ(w x)"]
    for token in neural_tokens:
        assert token not in bs_ppt_text, f"Topic B PPTX leaked Topic A token '{token}'"
        assert token not in bs_svg_text, f"Topic B SVG leaked Topic A token '{token}'"


def test_profile_isolation_and_pedagogical_differentiation(tmp_path):
    """55. Profile isolation: Student A (strength) vs Student B (missing knowledge) on same topic/evidence."""
    prod_profile_path = Path("data/profiles/student-001.json")
    prod_bytes_before = prod_profile_path.read_bytes()

    ai = build_system(str(tmp_path), test_mode=True)
    student_a = ai.create_or_load_student("student-A", "Alice", "Computer Science", "en")
    student_a.concept_mastery = {"Binary search": 1.0}
    student_a.strengths = ["Binary search"]
    student_a.unknown_concepts = []
    student_a.weak_concepts = []
    ai.store.save(student_a)

    student_b = ai.create_or_load_student("student-B", "Bob", "Computer Science", "en")
    student_b.concept_mastery = {"Binary search": 0.0}
    student_b.strengths = []
    student_b.unknown_concepts = ["Binary search"]
    student_b.weak_concepts = []
    ai.store.save(student_b)

    topic = "Binary search"
    evidence = [
        Evidence(
            chunk_id="c_iso",
            authority="university",
            text="Binary search requires an array sorted in ascending order. It checks the middle element. Worst-case runtime is O(log n). Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]

    out_a, routes_a = ai.generate_content("student-A", topic, ["explanation", "practice", "flashcards"], "en")
    out_b, routes_b = ai.generate_content("student-B", topic, ["explanation", "practice", "flashcards"], "en")

    # Scaffolding differentiation: Student A gets advanced, Student B gets foundational
    assert out_a["explanation"]["strategy"] == "higher-difficulty"
    assert out_a["explanation"]["learner_state"] == "strength"
    assert "Advanced" in out_a["explanation"]["content"] or "mastery" in out_a["explanation"]["content"].lower()

    assert out_b["explanation"]["strategy"] == "foundational"
    assert out_b["explanation"]["learner_state"] == "missing_knowledge"
    assert "Foundational" in out_b["explanation"]["content"] or "basics" in out_b["explanation"]["content"].lower()

    # Profile isolation: profiles do not cross-contaminate
    refreshed_a = ai.get_student_profile("student-A")
    refreshed_b = ai.get_student_profile("student-B")
    assert refreshed_a.strengths == ["Binary search"]
    assert refreshed_a.unknown_concepts == []
    assert refreshed_b.strengths == []
    assert refreshed_b.unknown_concepts == ["Binary search"]
    assert len(refreshed_a.generated_resources) == 1
    assert len(refreshed_b.generated_resources) == 1

    # Production profile byte integrity strictly preserved
    prod_bytes_after = prod_profile_path.read_bytes()
    assert prod_bytes_before == prod_bytes_after


def test_flashcard_typed_ir_and_semantic_relevance_validation(tmp_path):
    """56. Flashcards: Typed FlashcardItem, validate Back directly answers Front, no duplicate cards, fewer complete cards allowed."""
    from src.content_generation.service import FlashcardItem, ResourceContractValidator

    ev = [
        Evidence(
            chunk_id="c_fc_sem",
            authority="university",
            text="Binary search requires an array sorted in ascending order. It checks the median element. If the target equals the median element, its index is returned immediately. Worst-case runtime is O(log n). Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]

    # Verify typed FlashcardItem instantiation
    item = FlashcardItem(front="What is worst-case runtime?", back="Worst-case runtime is O(log n).", concept="runtime", evidence_ref="c_fc_sem")
    assert item.front == "What is worst-case runtime?"
    assert item.back == "Worst-case runtime is O(log n)."

    # Mismatched cards where Back does NOT answer Front
    mismatched_text = (
        "# Flashcards: Binary search\n\n"
        "**Card 1**\n"
        "**Front**: What is the worst-case runtime bound for binary search?\n"
        "**Back**: If the target equals the median element, its index is returned immediately.\n\n"
        "**Card 2**\n"
        "**Front**: What space complexity or memory overhead does binary search require?\n"
        "**Back**: Iterative space is O(1).\n\n"
        "**Card 3**\n"
        "**Front**: What is the primary operational requirement of binary search?\n"
        "**Back**: Binary search requires an array sorted in ascending order.\n"
    )
    is_valid, normalized, errors = ResourceContractValidator.validate_and_normalize(
        "flashcards", mismatched_text, ev, "Binary search"
    )
    assert not is_valid
    assert any("semantic mismatch" in err.lower() for err in errors)
    # Card 1 was dropped due to semantic mismatch, but Cards 2 and 3 survived and were normalized
    assert "Iterative space is O(1)" in normalized
    assert "sorted in ascending order" in normalized
    assert "median element, its index is returned immediately" not in normalized

    # Deduplication test
    dup_text = (
        "# Flashcards: Binary search\n\n"
        "**Card 1**\n"
        "**Front**: What is the worst-case runtime bound for binary search?\n"
        "**Back**: Worst-case runtime is O(log n).\n\n"
        "**Card 2**\n"
        "**Front**: What is the worst-case runtime bound for binary search?\n"
        "**Back**: Worst-case runtime is O(log n).\n"
    )
    is_v, norm_dup, err_dup = ResourceContractValidator.validate_and_normalize("flashcards", dup_text, ev, "Binary search")
    assert not is_v
    assert norm_dup.count("**Card 1**") == 1
    assert "**Card 2**" not in norm_dup


def test_formal_exam_typed_ir_and_synchronized_answer_key(tmp_path):
    """57. Formal Exam: Typed ExamQuestion, complete MCQ options (A/B/C/D), Short Answer, Analytical, Applied, synchronized Answer Key."""
    from src.content_generation.service import ExamQuestion, ResourceContractValidator

    ev = [
        Evidence(
            chunk_id="c_ex_sem",
            authority="university",
            text="Binary search requires an array sorted in ascending order. It checks the middle element. Worst-case runtime is O(log n). Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]

    eq = ExamQuestion(
        section="Section A: Multiple Choice",
        question_type="mcq",
        prompt="What is the runtime bound?",
        options=["A) O(1)", "B) O(n)", "C) O(log n)", "D) O(n^2)"],
        correct_answer="C",
        rubric="Section A: C (2 pts)",
        points=2,
    )
    assert len(eq.options) == 4

    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    res = gen.generate("exam", "Binary search", "en", evidence=ev, source_type="upload")
    assert res["type"] == "text"
    content = res["content"]

    # Verify Section A has 4 complete options
    assert "## Section A: Multiple Choice" in content
    assert "A)" in content and "B)" in content and "C)" in content and "D)" in content
    # Verify Short Answer, Analytical, and Applied sections
    assert "## Section B: Short Answer" in content
    assert "## Section C: Analytical Derivation" in content
    assert "## Section D: Practical Implementation" in content
    # Verify Answer Key & Scoring Rubric synchronized with questions
    assert "### Answer Key & Scoring Rubric" in content
    assert "Section A:" in content and "Section B:" in content


def test_analogy_strict_upload_mode_clean_refusal_when_no_analogy(tmp_path):
    """58. Analogy: Strict Selected-Upload Mode returns clean refusal when evidence lacks analogy, no phone-book invention."""
    ev = [
        Evidence(
            chunk_id="c_an_sem",
            authority="university",
            text="Binary search requires an array sorted in ascending order. It checks the middle element. Worst-case runtime is O(log n). Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    res = gen.generate("analogy", "Binary search", "en", evidence=ev, source_type="upload")
    assert res["type"] == "text"
    content = res["content"]

    # Verify clean evidence-limited refusal
    assert "# Grounded Analogy:" in content
    assert "The selected material does not contain an evidence-grounded analogy for this topic." in content
    # Verify zero phone-book / telephone directory invention
    assert "telephone directory" not in content.lower()
    assert "phone book" not in content.lower()
    assert "telephone" not in content.lower()
    # Verify 0 unsupported claims
    assert ContentGenerator._find_unsupported_claims(content, ev, strict_upload=True) == []


def test_question_bank_typed_ir_and_paired_qa_reliability(tmp_path):
    """59. Question Bank: Typed QuestionBankItem, paired Q&A, questions not treated as claims, never temporarily unavailable."""
    from src.content_generation.service import QuestionBankItem

    ev = [
        Evidence(
            chunk_id="c_qb_sem",
            authority="university",
            text="Binary search requires an array sorted in ascending order. It checks the middle element. Worst-case runtime is O(log n). Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]

    # Verify typed QuestionBankItem instantiation
    qb_item = QuestionBankItem(
        bloom_level="Remember",
        question="State the worst-case runtime complexity of Binary search.",
        reference_answer="Worst-case runtime is O(log n).",
        evidence_refs=["c_qb_sem"],
        learner_strategy="foundational",
    )
    assert qb_item.bloom_level == "Remember"
    assert "*Answer*:" not in qb_item.reference_answer

    # Verify generation reliability (must NOT return unavailable or raise exception)
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    res = gen.generate("question_bank", "Binary search", "en", evidence=ev, source_type="upload")
    assert res["type"] == "text"
    assert res.get("type") != "unavailable"
    content = res["content"]

    # Verify Bloom tags and paired Q&A
    assert "[Remember]" in content or "[Understand]" in content
    assert "*Answer*:" in content
    # Verify 0 unsupported claims
    assert ContentGenerator._find_unsupported_claims(content, ev, strict_upload=True) == []


# ==============================================================================
# THREE-RESOURCE SEMANTIC CLOSURE REGRESSION TESTS (1 - 17)
# ==============================================================================


def test_flashcard_back_rejects_raw_metadata_header():
    """1. Flashcard back rejects raw metadata header as answer."""
    ev = [
        Evidence(
            chunk_id="c1",
            authority="prof",
            text="Course: CS101 Introduction to Algorithms & Systems Topic: Divide-and-Conquer Search and Computational Limits Core Theoretical Foundations: 1. If target is smaller, search domain is restricted to lower half.",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.1,
        )
    ]
    raw = (
        "# Flashcards: Binary Search\n\n"
        "**Card 1**\n"
        "**Front**: How does binary search process or partition the search space?\n"
        "**Back**: Course: CS101 Introduction to Algorithms & Systems Topic: Divide-and-Conquer Search and Computational Limits Core Theoretical Foundations: 1.\n\n"
        "**Card 2**\n"
        "**Front**: What is the worst-case runtime complexity?\n"
        "**Back**: Worst-case runtime is O(log n).\n\n"
        "**Card 3**\n"
        "**Front**: What space complexity does iterative search require?\n"
        "**Back**: Space complexity is O(1).\n"
    )
    is_v, _, errors = ResourceContractValidator.validate_and_normalize("flashcards", raw, ev, "Binary Search")
    assert not is_v
    assert any("raw metadata" in e.lower() for e in errors)


def test_flashcard_back_with_combined_time_space_fails_when_front_asks_runtime_only():
    """2. Flashcard back with combined time+space fails when front asks runtime only."""
    ev = [
        Evidence(
            chunk_id="c1",
            authority="prof",
            text="The worst-case time complexity is O(log n), and space complexity is O(1) for iterative implementations. Precondition: array must be sorted. Halving: domain restricted to lower or upper half.",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.1,
        )
    ]
    raw = (
        "# Flashcards: Binary Search\n\n"
        "**Card 1**\n"
        "**Front**: What is the worst-case runtime complexity bound for binary search?\n"
        "**Back**: The worst-case time complexity is O(log n), and space complexity is O(1) for iterative implementations.\n\n"
        "**Card 2**\n"
        "**Front**: How does binary search partition space?\n"
        "**Back**: It restricts search to lower or upper half.\n\n"
        "**Card 3**\n"
        "**Front**: What precondition is required?\n"
        "**Back**: The array must be sorted in order.\n"
    )
    is_v, _, errors = ResourceContractValidator.validate_and_normalize("flashcards", raw, ev, "Binary Search")
    assert not is_v
    assert any("combined answer" in e.lower() or "runtime" in e.lower() for e in errors)


def test_flashcard_back_with_combined_time_space_fails_when_front_asks_space_only():
    """3. Flashcard back with combined time+space fails when front asks space only."""
    ev = [
        Evidence(
            chunk_id="c1",
            authority="prof",
            text="The worst-case time complexity is O(log n), and space complexity is O(1) for iterative implementations. Precondition: array must be sorted. Halving: domain restricted to lower or upper half.",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.1,
        )
    ]
    raw = (
        "# Flashcards: Binary Search\n\n"
        "**Card 1**\n"
        "**Front**: What space complexity or memory overhead does iterative binary search require?\n"
        "**Back**: The worst-case time complexity is O(log n), and space complexity is O(1) for iterative implementations.\n\n"
        "**Card 2**\n"
        "**Front**: How does binary search partition space?\n"
        "**Back**: It restricts search to lower or upper half.\n\n"
        "**Card 3**\n"
        "**Front**: What precondition is required?\n"
        "**Back**: The array must be sorted in order.\n"
    )
    is_v, _, errors = ResourceContractValidator.validate_and_normalize("flashcards", raw, ev, "Binary Search")
    assert not is_v
    assert any("combined answer" in e.lower() or "space" in e.lower() for e in errors)


def test_flashcard_duplicate_runtime_cards_rejected():
    """4. Flashcard duplicate runtime cards rejected."""
    ev = [
        Evidence(
            chunk_id="c1",
            authority="prof",
            text="The worst-case time complexity is O(log n), and space complexity is O(1) for iterative implementations. Precondition: array must be sorted. Halving: domain restricted to lower or upper half.",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.1,
        )
    ]
    raw = (
        "# Flashcards: Binary Search\n\n"
        "**Card 1**\n"
        "**Front**: What is the worst-case runtime complexity bound for binary search?\n"
        "**Back**: The worst-case time complexity is O(log n).\n\n"
        "**Card 2**\n"
        "**Front**: What mathematical asymptotic runtime governs binary search?\n"
        "**Back**: The worst-case time complexity is O(log n).\n\n"
        "**Card 3**\n"
        "**Front**: What space complexity does it use?\n"
        "**Back**: Space complexity is O(1).\n"
    )
    is_v, _, errors = ResourceContractValidator.validate_and_normalize("flashcards", raw, ev, "Binary Search")
    assert not is_v
    assert any("duplicate runtime" in e.lower() or "semantic duplicate" in e.lower() for e in errors)


def test_flashcard_prerequisite_question_requires_sorted_requirement_in_back():
    """5. Flashcard prerequisite question requires sorted requirement in back."""
    ev = [
        Evidence(
            chunk_id="c1",
            authority="prof",
            text="The indispensable precondition before binary search can be applied is that the data sequence must be strictly sorted. Worst-case is O(log n).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.1,
        )
    ]
    raw = (
        "# Flashcards: Binary Search\n\n"
        "**Card 1**\n"
        "**Front**: What fundamental input precondition is required before applying binary search?\n"
        "**Back**: Binary search operates in O(log n) time complexity.\n\n"
        "**Card 2**\n"
        "**Front**: What is the worst-case time complexity?\n"
        "**Back**: Worst-case is O(log n).\n\n"
        "**Card 3**\n"
        "**Front**: What space complexity is used?\n"
        "**Back**: Space is O(1).\n"
    )
    is_v, _, errors = ResourceContractValidator.validate_and_normalize("flashcards", raw, ev, "Binary Search")
    assert not is_v
    assert any("sorted" in e.lower() or "prerequisite" in e.lower() for e in errors)


def test_flashcard_partition_question_requires_lower_upper_half_behavior_in_back():
    """6. Flashcard partition question requires lower/upper-half behavior in back."""
    ev = [
        Evidence(
            chunk_id="c1",
            authority="prof",
            text="If the target is smaller, the search domain is restricted to the lower half; if larger, to the upper half. Runtime is O(log n). Space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.1,
        )
    ]
    raw = (
        "# Flashcards: Binary Search\n\n"
        "**Card 1**\n"
        "**Front**: How does binary search process or partition the search space?\n"
        "**Back**: Binary search checks items in the array.\n\n"
        "**Card 2**\n"
        "**Front**: What is the worst-case time complexity?\n"
        "**Back**: Worst-case is O(log n).\n\n"
        "**Card 3**\n"
        "**Front**: What space complexity is used?\n"
        "**Back**: Space is O(1).\n"
    )
    is_v, _, errors = ResourceContractValidator.validate_and_normalize("flashcards", raw, ev, "Binary Search")
    assert not is_v
    assert any("partition" in e.lower() or "half" in e.lower() for e in errors)


def test_exam_question_requires_multiple_answer_facets():
    """7. Exam question requires multiple answer facets."""
    q = ExamQuestion(
        section="Section B: Short-Answer Analysis",
        question_type="short_answer",
        prompt="Explain why binary search achieves O(log n) worst-case time complexity and state the indispensable data precondition required for this bound.",
        options=[],
        correct_answer="Binary search repeatedly halves the search space achieving O(log n) time, and requires a strictly sorted array as its precondition.",
        rubric="2 pts: O(log n) halving mechanism; 1 pt: sorted precondition requirement.",
        points=3,
        required_answer_facets=["time", "sorted"],
    )
    assert len(q.required_answer_facets) == 2
    assert "time" in q.required_answer_facets
    assert "sorted" in q.required_answer_facets


def test_exam_question_rejected_when_answer_only_covers_subset_of_required_facets():
    """8. Exam question rejected when answer only covers subset of required facets."""
    ok, err = ResourceContractValidator._validate_exam_question_facets(
        "Explain why binary search achieves O(log n) worst-case time complexity and state the indispensable data precondition required for this bound.",
        "Binary search achieves O(log n) time complexity.",
    )
    assert not ok
    assert "missing facet" in err.lower() or "sorted" in err.lower()


def test_exam_rejected_when_section_has_heading_but_zero_questions():
    """9. Exam rejected when section has heading but zero questions."""
    ev = [
        Evidence(
            chunk_id="c1",
            authority="prof",
            text="Binary search requires a sorted array. Worst-case runtime is O(log n). Iterative space is O(1). If target smaller, lower half; if larger, upper half.",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.1,
        )
    ]
    raw_exam = (
        "# Formal Examination: Binary Search\n\n"
        "## Section A: Multiple-Choice Diagnostic\n"
        "### Question 1.1\n"
        "What is worst-case time complexity?\n"
        "A) O(1)\nB) O(log n)\nC) O(n)\nD) O(n^2)\n\n"
        "## Section B: Short-Answer Analysis\n"
        "### Question 2.1\n"
        "What precondition is required for binary search to operate correctly?\n\n"
        "## Section C: Analytical Derivation\n"
        "### Question 3.1\n"
        "Explain how domain halving leads to logarithmic time and state iterative space complexity.\n\n"
        "## Section D: Practical Implementation & Trace Analysis\n\n"
        "## Answer Key\n"
        "1.1: B) O(log n)\n"
        "2.1: The indispensable precondition before binary search can be applied is that the data sequence must be strictly sorted.\n"
        "3.1: Binary search repeatedly halves search domain yielding worst-case time complexity O(log n), and space complexity is O(1) for iterative implementations.\n\n"
        "## Scoring Rubric\n"
        "- 1.1: 1 pt for selecting B.\n"
        "- 2.1: 2 pts for identifying sorted precondition.\n"
        "- 3.1: 3 pts (2 pts for halving domain yielding O(log n), 1 pt for iterative O(1) space complexity).\n"
    )
    is_v, _, errors = ResourceContractValidator.validate_and_normalize("exam", raw_exam, ev, "Binary Search")
    assert not is_v
    assert any("empty section" in e.lower() or "orphan" in e.lower() or "zero questions" in e.lower() for e in errors)


def test_exam_rejected_when_answer_key_has_fewer_entries_than_total_questions():
    """10. Exam rejected when Answer Key has fewer entries than total questions."""
    ev = [
        Evidence(
            chunk_id="c1",
            authority="prof",
            text="Binary search requires a sorted array. Worst-case runtime is O(log n). Iterative space is O(1). If target smaller, lower half; if larger, upper half.",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.1,
        )
    ]
    raw_exam = (
        "# Formal Examination: Binary Search\n\n"
        "## Section A: Multiple-Choice Diagnostic\n"
        "### Question 1.1\n"
        "What is worst-case time complexity?\n"
        "A) O(1)\nB) O(log n)\nC) O(n)\nD) O(n^2)\n\n"
        "## Section B: Short-Answer Analysis\n"
        "### Question 2.1\n"
        "What precondition is required for binary search to operate correctly?\n\n"
        "## Section C: Analytical Derivation\n"
        "### Question 3.1\n"
        "Explain how domain halving leads to logarithmic time and state iterative space complexity.\n\n"
        "## Section D: Practical Implementation & Trace Analysis\n"
        "### Question 4.1\n"
        "Trace binary search branching behavior when target is smaller and larger.\n\n"
        "## Answer Key\n"
        "1.1: B) O(log n)\n"
        "2.1: The indispensable precondition before binary search can be applied is that the data sequence must be strictly sorted.\n\n"
        "## Scoring Rubric\n"
        "- 1.1: 1 pt for selecting B.\n"
        "- 2.1: 2 pts for identifying sorted precondition.\n"
        "- 3.1: 3 pts (2 pts for halving domain yielding O(log n), 1 pt for iterative O(1) space complexity).\n"
        "- 4.1: 3 pts (1.5 pts for lower-half branch when smaller, 1.5 pts for upper-half branch when larger).\n"
    )
    is_v, _, errors = ResourceContractValidator.validate_and_normalize("exam", raw_exam, ev, "Binary Search")
    assert not is_v
    assert any("answer key" in e.lower() for e in errors)


def test_exam_rejected_when_scoring_rubric_has_fewer_entries_than_total_questions():
    """11. Exam rejected when Scoring Rubric has fewer entries than total questions."""
    ev = [
        Evidence(
            chunk_id="c1",
            authority="prof",
            text="Binary search requires a sorted array. Worst-case runtime is O(log n). Iterative space is O(1). If target smaller, lower half; if larger, upper half.",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.1,
        )
    ]
    raw_exam = (
        "# Formal Examination: Binary Search\n\n"
        "## Section A: Multiple-Choice Diagnostic\n"
        "### Question 1.1\n"
        "What is worst-case time complexity?\n"
        "A) O(1)\nB) O(log n)\nC) O(n)\nD) O(n^2)\n\n"
        "## Section B: Short-Answer Analysis\n"
        "### Question 2.1\n"
        "What precondition is required for binary search to operate correctly?\n\n"
        "## Section C: Analytical Derivation\n"
        "### Question 3.1\n"
        "Explain how domain halving leads to logarithmic time and state iterative space complexity.\n\n"
        "## Section D: Practical Implementation & Trace Analysis\n"
        "### Question 4.1\n"
        "Trace binary search branching behavior when target is smaller and larger.\n\n"
        "## Answer Key\n"
        "1.1: B) O(log n)\n"
        "2.1: The indispensable precondition before binary search can be applied is that the data sequence must be strictly sorted.\n"
        "3.1: Binary search repeatedly halves search domain yielding worst-case time complexity O(log n), and space complexity is O(1) for iterative implementations.\n"
        "4.1: If target is smaller, search domain is restricted to the lower half; if larger, to the upper half.\n\n"
        "## Scoring Rubric\n"
        "- 1.1: 1 pt for selecting B.\n"
        "- 2.1: 2 pts for identifying sorted precondition.\n"
    )
    is_v, _, errors = ResourceContractValidator.validate_and_normalize("exam", raw_exam, ev, "Binary Search")
    assert not is_v
    assert any("scoring rubric" in e.lower() or "rubric" in e.lower() for e in errors)


def test_question_bank_question_premise_rejected_when_target_not_present_not_in_evidence():
    """12. Question Bank question premise rejected when target-not-present is not in evidence."""
    ev = [
        Evidence(
            chunk_id="c1",
            authority="prof",
            text="Binary search checks the middle element. If equal, return immediately. If smaller, search lower half; if larger, upper half. Worst-case is O(log n).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.1,
        )
    ]
    raw_qb = (
        "# Question Bank: Binary Search\n\n"
        "1. [Understand] Under what condition does binary search terminate when a target value is not present in the array?\n"
        "   *Answer*: It stops when pointers cross.\n\n"
        "2. [Remember] What is the worst-case time complexity of binary search?\n"
        "   *Answer*: Worst-case is O(log n).\n"
    )
    is_v, _, errors = ResourceContractValidator.validate_and_normalize("question_bank", raw_qb, ev, "Binary Search")
    assert not is_v
    assert any("unsupported" in e.lower() or "target" in e.lower() for e in errors)


def test_question_bank_reference_answer_rejected_when_pointer_crossing_not_in_evidence():
    """13. Question Bank reference answer rejected when pointer-crossing concept is not in evidence."""
    ev = [
        Evidence(
            chunk_id="c1",
            authority="prof",
            text="Binary search checks the middle element. Worst-case is O(log n). Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.1,
        )
    ]
    raw_qb = (
        "# Question Bank: Binary Search\n\n"
        "1. [Understand] How does binary search partition the search domain?\n"
        "   *Answer*: It adjusts low and high pointers until low exceeds high when pointer crossing occurs.\n\n"
        "2. [Remember] What is the worst-case time complexity of binary search?\n"
        "   *Answer*: Worst-case is O(log n).\n"
    )
    is_v, _, errors = ResourceContractValidator.validate_and_normalize("question_bank", raw_qb, ev, "Binary Search")
    assert not is_v
    assert any("pointer crossing" in e.lower() or "unsupported" in e.lower() for e in errors)


def test_question_bank_item_rejected_when_question_contains_raw_metadata_header():
    """14. Question Bank item rejected when question contains raw metadata header."""
    ev = [
        Evidence(
            chunk_id="c1",
            authority="prof",
            text="Course: CS101 Introduction to Algorithms & Systems Topic: Divide-and-Conquer Search 1. Worst-case is O(log n). Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.1,
        )
    ]
    raw_qb = (
        "# Question Bank: Binary Search\n\n"
        "1. [Remember] Course: CS101 Introduction to Algorithms & Systems Topic: Divide-and-Conquer Search: What is the runtime?\n"
        "   *Answer*: Worst-case is O(log n).\n\n"
        "2. [Understand] What is the space complexity?\n"
        "   *Answer*: Space complexity is O(1).\n"
    )
    is_v, _, errors = ResourceContractValidator.validate_and_normalize("question_bank", raw_qb, ev, "Binary Search")
    assert not is_v
    assert any("raw metadata" in e.lower() for e in errors)


def test_question_bank_item_rejected_when_reference_answer_contains_raw_metadata_header():
    """15. Question Bank item rejected when reference answer contains raw metadata header."""
    ev = [
        Evidence(
            chunk_id="c1",
            authority="prof",
            text="Course: CS101 Introduction to Algorithms & Systems Topic: Divide-and-Conquer Search 1. Worst-case is O(log n). Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.1,
        )
    ]
    raw_qb = (
        "# Question Bank: Binary Search\n\n"
        "1. [Remember] What is the worst-case runtime complexity?\n"
        "   *Answer*: Course: CS101 Introduction to Algorithms & Systems Topic: Divide-and-Conquer Search: Worst-case is O(log n).\n\n"
        "2. [Understand] What is the space complexity?\n"
        "   *Answer*: Space complexity is O(1).\n"
    )
    is_v, _, errors = ResourceContractValidator.validate_and_normalize("question_bank", raw_qb, ev, "Binary Search")
    assert not is_v
    assert any("raw metadata" in e.lower() for e in errors)


def test_question_bank_allows_fewer_complete_grounded_items_when_evidence_limited():
    """16. Question Bank allows fewer complete grounded items when evidence is limited."""
    ev = [
        Evidence(
            chunk_id="c1",
            authority="prof",
            text="Binary search requires an array sorted in ascending order. Worst-case runtime is O(log n). Iterative space is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=0.9,
            score=0.1,
        )
    ]
    raw_qb = (
        "# Question Bank: Binary Search\n\n"
        "1. [Remember] What fundamental precondition must hold before binary search can execute?\n"
        "   *Answer*: The array must be sorted in ascending order.\n\n"
        "2. [Understand] What is the worst-case runtime complexity of binary search?\n"
        "   *Answer*: Worst-case runtime is O(log n).\n"
    )
    is_v, norm, errors = ResourceContractValidator.validate_and_normalize("question_bank", raw_qb, ev, "Binary Search")
    assert is_v
    assert len(errors) == 0
    assert "1." in norm and "2." in norm


def test_full_run_on_sample_manual_lecture_generates_valid_tools_without_defects(tmp_path):
    """17. Full run on sample_manual_lecture.txt generates valid Flashcards, Formal Exam, and Question Bank without 'temporarily unavailable' and without metadata/scope/rubric defects."""
    with open("data/manual_test/sample_manual_lecture.txt") as f:
        lecture_text = f.read()

    ev = [
        Evidence(
            chunk_id="c_manual",
            authority="prof",
            text=lecture_text,
            source="sample_manual_lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))

    # 1. Flashcards
    fc_res = gen.generate("flashcards", "Binary search", "en", evidence=ev, source_type="student_upload")
    assert fc_res["type"] == "text"
    assert "temporarily unavailable" not in fc_res["content"].lower()
    assert "Course:" not in fc_res["content"]
    assert "Core Theoretical Foundations:" not in fc_res["content"]
    assert "O(1)" in fc_res["content"]
    assert "o(1)" not in fc_res["content"]
    cards = re.split(r"(?m)^\*\*Card\s*\d+\*\*", fc_res["content"])[1:]
    fronts = [re.search(r"\*\*Front\*\*:\s*(.*?)(?=\n\s*\*\*Back\*\*|\Z)", c, re.S).group(1).lower() for c in cards if "**Front**:" in c]
    runtime_cards = [f for f in fronts if "runtime" in f or "time complexity" in f]
    assert len(runtime_cards) == 1, f"Found multiple runtime cards: {runtime_cards}"

    # 2. Formal Exam
    exam_res = gen.generate("exam", "Binary search", "en", evidence=ev, source_type="student_upload")
    assert exam_res["type"] == "text"
    assert "temporarily unavailable" not in exam_res["content"].lower()
    assert "## Section A:" in exam_res["content"]
    assert "## Section B:" in exam_res["content"]
    assert "## Section C:" in exam_res["content"]
    assert "## Section D:" in exam_res["content"]
    assert "Answer Key" in exam_res["content"]
    assert "O(1)" in exam_res["content"]
    assert "o(1)" not in exam_res["content"]

    # 3. Question Bank
    qb_res = gen.generate("question_bank", "Binary search", "en", evidence=ev, source_type="student_upload")
    assert qb_res["type"] == "text"
    assert "temporarily unavailable" not in qb_res["content"].lower()
    assert "Course:" not in qb_res["content"]
    assert "Core Theoretical Foundations:" not in qb_res["content"]
    assert "pointer crossing" not in qb_res["content"].lower()
    assert "not found" not in qb_res["content"].lower()
    assert "low > high" not in qb_res["content"].lower()
    assert "*Answer*:" in qb_res["content"]
    assert "O(1)" in qb_res["content"]
    assert "o(1)" not in qb_res["content"]


# ==============================================================================
# ASYMPTOTIC NOTATION FIDELITY REGRESSION TESTS (1 - 7)
# ==============================================================================


def test_asymptotic_notation_o1_evidence_rejects_generated_little_o1():
    """Asymptotic fidelity 1: O(1) evidence rejects generated o(1)."""
    from src.content_generation.service import GenericClaimValidator
    ev = [
        Evidence(
            chunk_id="c_asymp1",
            authority="university",
            text="Iterative space complexity is O(1) for binary search. Worst-case runtime is O(log n).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    # Valid claim with exact Big-O
    assert GenericClaimValidator.classify_claim("Iterative space complexity is O(1).", ev) == "SUPPORTED"
    # Invalid claim with little-o must be rejected
    assert GenericClaimValidator.classify_claim("Iterative space complexity is o(1).", ev) == "UNSUPPORTED"
    ok, err = GenericClaimValidator.validate_asymptotic_fidelity("Iterative space complexity is o(1).", ev)
    assert not ok
    assert "o(1)" in err and "O(1)" in err


def test_asymptotic_notation_ologn_evidence_rejects_generated_little_ologn():
    """Asymptotic fidelity 2: O(log n) evidence rejects o(log n)."""
    from src.content_generation.service import GenericClaimValidator
    ev = [
        Evidence(
            chunk_id="c_asymp2",
            authority="university",
            text="The worst-case runtime complexity is O(log n). Space complexity is O(1).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    # Valid claim with exact Big-O
    assert GenericClaimValidator.classify_claim("The worst-case runtime complexity is O(log n).", ev) == "SUPPORTED"
    # Invalid claim with little-o must be rejected
    assert GenericClaimValidator.classify_claim("The worst-case runtime complexity is o(log n).", ev) == "UNSUPPORTED"
    ok, err = GenericClaimValidator.validate_asymptotic_fidelity("The worst-case runtime complexity is o(log n).", ev)
    assert not ok
    assert "o(log n)" in err and "O(log n)" in err


def test_asymptotic_notation_legitimate_little_o_evidence_preserved():
    """Asymptotic fidelity 3: Legitimate little-o evidence, if explicitly supplied, is NOT converted to Big-O."""
    from src.content_generation.service import GenericClaimValidator
    ev = [
        Evidence(
            chunk_id="c_asymp3",
            authority="university",
            text="The asymptotic remainder error is strictly bounded by o(n) in the limit.",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    # Legitimate little-o in evidence must be accepted
    assert GenericClaimValidator.classify_claim("The asymptotic remainder error is strictly bounded by o(n).", ev) == "SUPPORTED"
    # Big-O must NOT be accepted when evidence only specifies little-o
    assert GenericClaimValidator.classify_claim("The asymptotic remainder error is strictly bounded by O(n).", ev) == "UNSUPPORTED"
    # Repairing little-o must NOT convert it to Big-O
    repaired = GenericClaimValidator.repair_asymptotic_notation("The remainder error is o(n).", ev)
    assert "o(n)" in repaired and "O(n)" not in repaired


def test_asymptotic_notation_greek_symbols_remain_distinct():
    """Asymptotic fidelity 4: Greek asymptotic symbols (Ω, ω, Θ) remain distinct."""
    from src.content_generation.service import GenericClaimValidator
    ev = [
        Evidence(
            chunk_id="c_asymp4",
            authority="university",
            text="The algorithmic lower bound is Ω(n), tight bound is Θ(n), and upper bound is O(n).",
            source="lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    # Supported Greek tokens
    assert GenericClaimValidator.classify_claim("The algorithmic lower bound is Ω(n).", ev) == "SUPPORTED"
    assert GenericClaimValidator.classify_claim("The algorithmic tight bound is Θ(n).", ev) == "SUPPORTED"
    assert GenericClaimValidator.classify_claim("The algorithmic upper bound is O(n).", ev) == "SUPPORTED"
    # Unsupported Greek token: little-omega ω(n) is absent from evidence
    assert GenericClaimValidator.classify_claim("The algorithmic lower bound is ω(n).", ev) == "UNSUPPORTED"
    ok, _ = GenericClaimValidator.validate_asymptotic_fidelity("The algorithmic lower bound is ω(n).", ev)
    assert not ok


def test_flashcard_preserves_capital_o1_complexity(tmp_path):
    """Asymptotic fidelity 5: Flashcard preserves exact O(1) from sample_manual_lecture.txt."""
    with open("data/manual_test/sample_manual_lecture.txt") as f:
        lecture_text = f.read()

    ev = [
        Evidence(
            chunk_id="c_manual",
            authority="prof",
            text=lecture_text,
            source="sample_manual_lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    res = gen.generate("flashcards", "Binary search", "en", evidence=ev, source_type="student_upload")
    assert res["type"] == "text"
    content = res["content"]
    assert "O(1)" in content
    assert "o(1)" not in content
    cards = re.split(r"(?m)^\*\*Card\s*\d+\*\*", content)[1:]
    space_cards = [c for c in cards if "space complexity" in c.lower() or "memory overhead" in c.lower()]
    assert len(space_cards) >= 1
    for sc in space_cards:
        assert "O(1)" in sc
        assert "o(1)" not in sc


def test_question_bank_preserves_capital_o1_complexity(tmp_path):
    """Asymptotic fidelity 6: Question Bank preserves exact O(1) from sample_manual_lecture.txt."""
    with open("data/manual_test/sample_manual_lecture.txt") as f:
        lecture_text = f.read()

    ev = [
        Evidence(
            chunk_id="c_manual",
            authority="prof",
            text=lecture_text,
            source="sample_manual_lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    res = gen.generate("question_bank", "Binary search", "en", evidence=ev, source_type="student_upload")
    assert res["type"] == "text"
    content = res["content"]
    assert "O(1)" in content
    assert "o(1)" not in content
    items = re.split(r"(?m)^\d+\.\s+", content)[1:]
    space_items = [it for it in items if "space complexity" in it.lower() or "memory overhead" in it.lower()]
    assert len(space_items) >= 1
    for si in space_items:
        assert "O(1)" in si
        assert "o(1)" not in si


def test_formal_exam_preserves_capital_o1_complexity(tmp_path):
    """Asymptotic fidelity 7: Formal Exam preserves exact O(1) from sample_manual_lecture.txt."""
    with open("data/manual_test/sample_manual_lecture.txt") as f:
        lecture_text = f.read()

    ev = [
        Evidence(
            chunk_id="c_manual",
            authority="prof",
            text=lecture_text,
            source="sample_manual_lecture.txt",
            source_type="student_upload",
            trust_score=1.0,
            score=0.1,
        )
    ]
    gen = ContentGenerator(artifact_dir=str(tmp_path / "artifacts"))
    res = gen.generate("exam", "Binary search", "en", evidence=ev, source_type="student_upload")
    assert res["type"] == "text"
    content = res["content"]
    assert "O(1)" in content
    assert "o(1)" not in content
    assert "## Section D: Practical Implementation" in content
    assert "O(1)" in content.split("## Section D: Practical Implementation")[1]





