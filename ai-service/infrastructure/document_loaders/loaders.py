from pathlib import Path
from src.core.exceptions import DocumentError

ALLOWED = {'.pdf', '.docx', '.pptx', '.txt', '.md'}

def safe_name(name: str) -> str:
    n = Path(name).name
    if n != name or Path(n).suffix.lower() not in ALLOWED:
        raise DocumentError('Unsafe or unsupported filename')
    return n

def extract(path: str | Path) -> list[tuple[int, str]]:
    p = Path(path)
    ext = p.suffix.lower()
    if ext in {'.txt', '.md'}:
        try:
            return [(1, p.read_text(encoding='utf-8', errors='ignore'))]
        except Exception as exc:
            raise DocumentError(f'Failed to read text file: {exc}') from exc

    if ext == '.pdf':
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(p))
            pages = [(i + 1, x.extract_text() or '') for i, x in enumerate(reader.pages)]
            total_text = ''.join(text for _, text in pages).strip()
            if not total_text:
                raise DocumentError('No readable text could be extracted from PDF (OCR is not configured)')
            return pages
        except DocumentError:
            raise
        except Exception as exc:
            raise DocumentError(f'Failed to parse PDF document: {exc}') from exc

    if ext == '.docx':
        try:
            from docx import Document
            doc = Document(str(p))
            parts = [x.text for x in doc.paragraphs if x.text]
            for table in doc.tables:
                for row in table.rows:
                    row_text = ' | '.join(cell.text.strip() for cell in row.cells)
                    if row_text:
                        parts.append(row_text)
            return [(1, '\n'.join(parts))]
        except Exception as exc:
            raise DocumentError(f'Failed to parse DOCX document: {exc}') from exc

    if ext == '.pptx':
        try:
            from pptx import Presentation
            prs = Presentation(str(p))
            slides = []
            for i, s in enumerate(prs.slides):
                slide_parts = []
                for sh in s.shapes:
                    if hasattr(sh, 'text') and sh.text:
                        slide_parts.append(sh.text)
                    if getattr(sh, 'has_table', False):
                        for row in sh.table.rows:
                            row_text = ' | '.join(cell.text.strip() for cell in row.cells)
                            if row_text:
                                slide_parts.append(row_text)
                slides.append((i + 1, '\n'.join(slide_parts)))
            return slides
        except Exception as exc:
            raise DocumentError(f'Failed to parse PPTX presentation: {exc}') from exc

    raise DocumentError(f'Unsupported file type: {ext}')
